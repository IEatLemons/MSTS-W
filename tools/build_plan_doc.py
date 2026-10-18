from pathlib import Path
from datetime import date
import json, csv, re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]
    s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    s.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Subtitle'].font.italic=False
doc.styles['Normal'].paragraph_format.line_spacing=1.14
for name,size in [('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(11 if name!='Title' else 0)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.core_properties.title='船舶管理财务系统建设方案与完整数据字典'
doc.core_properties.subject='模块功能 数据模型 计算公式 技术栈 局域网部署 备份与容灾'
doc.core_properties.author=''
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run('第 '); r.font.size=Pt(9)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)
foot.add_run(' 页').font.size=Pt(9)

def p(t='',style=None):
    return doc.add_paragraph(t.replace('。。','。'),style)
def h(t,level=1): return doc.add_heading(t,level)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=widths or ([17.6/len(headers)]*len(headers))
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for c,txt,w in zip(t.rows[0].cells,headers,widths): c.text=str(txt); c.width=Cm(w)
    for row in rows:
        cs=t.add_row().cells
        for c,txt,w in zip(cs,row,widths): c.text=str(txt); c.width=Cm(w)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); pr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            tcp.append(borders)
            marg=OxmlElement('w:tcMar')
            for edge,value in [('top',65),('bottom',65),('left',80),('right',80)]:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(value)); el.set(qn('w:type'),'dxa'); marg.append(el)
            tcp.append(marg)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E5EDF5' if i==0 else 'FFFFFF'); tcp.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(0); par.paragraph_format.space_before=Pt(0)
                par.paragraph_format.line_spacing=1.05
                for run in par.runs:
                    run.font.size=Pt(9)
                    if i==0: run.bold=True
    p().paragraph_format.space_after=Pt(0)
    return t
def link(title,url):
    par=p(); rel=par.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); prop=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'164D7A'); prop.append(color); run.append(prop)
    txt=OxmlElement('w:t'); txt.text=title; run.append(txt); hl.append(run); par._p.append(hl)

p('船舶管理财务系统建设方案', 'Title')
p('模块功能  数据字段  计算公式  技术架构  备份与容灾', 'Subtitle')
p('Vessel Management Finance System\n方案日期 2026年10月15日   版本 1.0')
p('本方案的目标是把目前分散在 Excel 中的费用、收款、银行和船长备用金，整理成一次录入、逐笔核销、按船查询、可以追溯的财务管理系统。首期在公司局域网部署，员工通过浏览器使用；数据库、凭证附件和备份由公司内部保存。')
p('建议先建设船舶资金与费用管理子账，同时建立凭证、期间和审计基础。正式总账及法定财务报表，应在财务确认会计科目、期初余额、资金性质和记账规则后启用。旧表 SOFP 作为历史兼容报表保留，不能直接当成已经符合会计准则的资产负债表。')
h('阅读导航',2)
table(['你想解决的问题','阅读位置'],[
('先看懂这个系统到底管什么','第1章 业务全景；第3章 完整业务例子'),
('需要开发哪些模块 每个模块做什么','第2章 模块规划'),
('模块中有哪些数据 数据如何关联','第4章 数据组织；第9章 完整数据字典'),
('金额和余额具体怎样计算','第5章 公式与口径；附录 原表规则逐项处置'),
('使用哪些技术 局域网如何部署','第6章 技术栈；第7章 备份与容灾'),
('如何分阶段落地和验收','第8章 实施计划和验收')],[7,10.6])
p('文中“现状”表示原工作簿及项目提取结果中的观察；“规划”表示建议建设的功能和数据模型；“待确认”表示需要由财务或业务负责人决定的政策。RPO 是最多可接受的数据丢失时长，RTO 是恢复服务所需的目标时间。文中的恢复时长均为设计目标，须通过演练验证。')

h('1 先用业务语言理解系统')
h('1.1 系统管理的五条主线',2)
table(['主线','要回答的问题','主要依据'],[
('费用','哪条船发生了什么费用 由谁承担 属于什么期间','订单 验收 发票 费用归属'),
('付款','公司或船上实际付了多少钱 还欠供应商多少','付款记录 付款核销 银行或现金流水'),
('船东结算','公司代船东付的钱 收回了多少 还应收多少','船东结算单 真实回款 回款核销'),
('资金','每个银行账户和船长现金账户还剩多少钱','账户期初 真实收支 对账单 现金盘点'),
('计划与监督','后续要多少钱 谁批准了 修改是否留痕','资金申请 审批 关账 审计 备份')],[2.4,7.8,7.4])
p('五个费用中心是船管、船东、接船、换旗和坞修。它们是同一套台账的分类维度，不需要建立五套相互复制的数据。每笔业务还要说明“资金由谁付、费用由谁承担、受益船是哪条”，这三个对象可以不同。')
h('1.2 当前项目已有的东西',2)
p('项目保存了1份原始工作簿、14个工作表、6926个非空单元格、2722条公式，以及业务记录、银行流水、科目和异常分析。独立复算中2721条与原表缓存一致，另1条复现原表错误。公式复算一致仅证明提取与计算一致，不能代替银行核销、凭证审核或会计政策确认。')
table(['原表或资料','业务理解','迁移处理'],[
('Expenses Lists 加五类 Cost 页','总清单和分类页存在重复视图','归并为唯一业务及分摊 不直接相加导入'),
('Shipmanagement RCVD','实际到账 结算缺口和计划拨款混在一起','实际收款与资金申请分开'),
('Bank Verification','22条流水 余额靠倒序行引用','导入稳定流水标识 按账户和日期计算'),
('CTM Expenses List','22条现金流水 含补给与消费','作为船长现金账户 补给和费用区分'),
('Owner 和 Take Over 等结算列','供应商付款状态与船东结算状态并存','各自核销 分别计算待付和待收'),
('分析 伙食海关库房管理 DD Cost','估算和库存金额尚未形成真实业务历史 坞修为空模板','保留结构 不伪造历史金额或库存'),
('附件列和科目页','附件多为文件名 两套科目共192条编码','附件补档 科目按命名空间整理')],[4.6,6.3,6.7])
h('1.3 读懂目前的几个金额',2)
p('下列数字来自原表保存于2026年10月14日的快照，单位为 USD，显示保留两位小数。原表没有完整按截止日筛选的机制，因此这里称为“保存快照”，不视为已审定的截至日账目。精确值及原公式保留在历史层。')
table(['指标','原表显示值','正确理解'],[
('五类已分配实收资金','371,554.52','不是全部营业收入 包含代管或往来性质 待财务认定'),
('旧表总费用','398,259.75','六状态分桶加手续费 包含混合业务 不是纯经营成本'),
('旧表资金减费用净额','-26,705.23','371,554.52 − 398,259.75 不是直接认定的亏损'),
('银行余额','92,779.85','需取得银行正式对账单确认'),
('船长备用金余额','29,908.00','船上现金资产 不是剩余可计费用'),
('已完成费用加手续费','251,657.25','来自旧表 Completed 状态 不等于已完成真实银行核销'),
('旧表资金平衡差额','-2,790.58','实收 − 银行 − CTM − 已完成费用含费 尚未解释'),
('现金覆盖差额','-23,914.65','银行 + CTM − 旧表五类未完成状态费用 不是利润')],[4.5,3.1,10])
p('旧表分类费用为：船管308,970.83，船东36,788.27，接船26,158.99，换旗26,341.66，坞修0。总额核对一致仍可能隐藏单笔金额或状态互换，因此上线前必须逐笔核对，不能仅凭总额一致认定迁移正确。')

MODULES=[
('M01','组织 船舶与基础资料','让每笔业务明确属于哪家公司 哪条船 何种币种和费用中心。','维护公司与船名别名；维护供应商 船东 员工等往来方；维护五类费用中心；维护两套业务科目和正式会计科目映射；维护带日期和方向的汇率。','公司 船舶 往来方 币种 汇率 费用中心 科目','organization vessel party currency fx_rate cost_center account_code gl_account','管理员和财务维护；其他模块只能选择有效主数据。重复名称合并保留别名，历史记录仍可追溯。','第一阶段'),
('M02','费用 采购与验收','说明费用从申请到验收是怎样形成的 并避免同笔业务重复录入。','建立唯一业务单；可无采购订单直接录入服务费；记录请购号 订单号 原币额；支持分批验收；按费用中心和受益船拆分；区分费用 押金 借款 跨船代付和转账；记录差异原因。','业务单 业务行 订单行 验收记录','business_event event_line purchase_order_line goods_receipt','业务录入申请与附件，采购补充订单，验收人员确认金额；订单额与验收额可以不同，但差异必须说明。','第一阶段'),
('M03','发票 应付与贷项','确定应该付给供应商多少钱 到期日是什么 还欠多少。','录入发票头和明细；匹配业务行及订单；保存发票原币及USD；处理贷项与部分抵扣；展示已付 未付及账龄；支持缺票应计并在到票后转回，避免双计。','发票 发票行 贷项抵扣 付款核销','invoice invoice_line credit_allocation payment_allocation','财务审核发票，系统依据核销计算待付；原表 C/N Profit 和 Profit 只保留历史字段，确认定义前不自动计算利润。','第一阶段'),
('M04','付款 退款与核销','把付款动作和实际资金流水对应起来 并支持部分支付。','批量付款申请；付款审批；人工网银执行后登记；一笔付款核销多张发票；一票分次付款；记录失败 原额退款 手续费返还；押金及借款付款关联对应业务。','付款单 发票核销 银行匹配 退款关联','payment payment_allocation bank_match','付款权限与审批权限分离。系统首期生成支付记录与核销，不直接接入银行支付；Completed 只能由实际支付和核销条件产生。','第一阶段'),
('M05','实际收款与资金分配','说明真实到账多少钱 每笔钱用于什么事项。','登记拨款 代付回款 借款归还等资金性质；分开保存毛额 手续费 净额；银行入账核实后分配五类资金池；支持分批分配和未分配余额；关联资金申请。','收款单 用途分配 银行匹配','receipt receipt_allocation bank_match','财务确认实际到账；计划拨款不进入本模块实收。用途分配不得超过净额，全部分配后未分配余额为零。','第一阶段'),
('M06','银行账户 流水与对账','确认账户真实余额 每条银行流水是否都有业务解释。','按账户导入对账单CSV或Excel；处理独立手续费 结息 退款；重复导入拦截；按稳定顺序计算余额；人工确认匹配建议；支持跨船资金归属和未核销项追踪。','银行账户 银行流水 匹配记录','bank_account bank_transaction bank_match','银行是资金事实来源；业务匹配不生成第二笔资金变动。账户币种分别汇总，不直接相加不同币种。','第一阶段'),
('M07','船长备用金与内部划拨','同时看清银行钱和船上现金 并保证补给不重复算费用。','维护按船按币种的现金账户；登记补给 消费 退回和调整；银行到现金或账户间划拨建立双端关联；现金消费核销同一发票；月末盘点并审批差异。','现金账户 现金流水 内部划拨','cash_account cash_movement internal_transfer','出纳登记补给，船上上传消费凭证；银行转入1000时银行减1000 现金增1000 经营费用不变。','第一阶段'),
('M08','船东结算与跨船往来','区分欠供应商的钱与应向船东或其他船收回的钱。','按期间选取可结算费用；区分公司代付和船东直接支付；形成结算单；分次回款核销；按受益船记录跨船代付；查看待收账龄及争议款。','结算单 结算行 结算回款核销 业务归属','owner_settlement owner_settlement_line owner_receipt_allocation event_line','供应商已收款不代表船东已报销。船东直接付款不增加公司现金支出；是否形成公司费用或往来由合同和会计政策决定。','第一阶段'),
('M09','资金计划 估算与补款申请','提前测算需要多少资金 并跟踪申请是否到账。','按日费率与天数生成计划；按资金缺口申请补款；支持批准金额和多次到账；维护低中高估算与合同分成；展示计划 批准 实收差异；参数按日期版本化。','资金申请 收款分配 估算项 合同参数','funding_request receipt_allocation estimate_item settlement_term','155000美元计划在到账前只影响预测。分析页中值算法及45%和55%分成归属未确认前，只允许人工审核的估算。','第一阶段资金申请 第二阶段复杂估算'),
('M10','伙食 招待品库存与月结','解释现金支出与实际耗用的区别 并连续结转每月余额。','登记金额型采购和耗用；保存期初与期末盘点估值；计算伙食和招待品月度收支；记录损耗和盘点差异；自动承接上月期末；关账后调整留痕。','库存月结 库存变动 现金关联','inventory_month inventory_movement','首期保留接口，取得盘点记录后上线。按金额管理；若未来要求每个物品的数量 批次和单价，需要另行扩展物料及仓库子模块。','第二阶段'),
('M11','报表 凭证与期间管理','按同一截止日展示资金 费用 应付 应收 并冻结已确认历史。','提供资金总览 五类费用 现金覆盖 应付账龄 船东待收；保留旧SOFP兼容视图；生成凭证草稿；审核后记账；借贷平衡校验；月结锁定；调整用冲销或更正；报表保存规则版本。','期间 凭证头 凭证明细 报表快照 指标','accounting_period journal_entry journal_line report_snapshot report_metric','管理报表可先上线；正式资产负债表和利润表须具备已确认期初余额 科目映射及会计政策。报表不通过修改股本或留存收益强行调平。','第一阶段子账和期间 正式总账按前置条件启用'),
('M12','权限 审批 附件 审计与历史迁移','让数据能被查证 控制谁能操作 并安全承接旧表。','用户与角色；按公司和船授权；审批分离；文件哈希和缺档占位；保留Excel哈希 工作表 单元格 原值和公式；重复记录人工归并；异常整改清单；敏感导出留痕。','用户 角色 授权 审批 附件 源文件 溯源 异常 审计','app_user role permission user_role role_permission workflow_action attachment attachment_link source_file source_cell_mapping migration_issue audit_log','导入先暂存再审核再入正式账；附件文件名不视为已归档；系统管理员不能直接修改已记账金额。','第一阶段'),
('M13','运维 备份与容灾','在误删 硬盘故障 服务器损坏和机房事故后能恢复业务。','持续监测数据库 附件 证书 磁盘和备份；记录备份完成及归档延迟；隔离备份账号；离线异地副本；季度恢复演练；保留事故与切换证据。','备份任务 恢复演练 独立备份仓库 清单与密钥','backup_run restore_drill','运维执行 财务验证 负责人批准恢复上线；数据库中的任务日志同时导出到独立仓库，生产数据库损坏后仍能查找备份。','第一阶段上线必备')
]
h('2 需要建设的模块与详细功能')
table(['模块','主要数据','建设顺序'],[(m[0]+' '+m[1],m[4],m[7]) for m in MODULES],[5.7,7,4.9])
for mid,title,purpose,features,data,tables,flow,phase in MODULES:
    h(mid+' '+title,2)
    p('作用：'+purpose)
    for i,s in enumerate(features.split('；'),1): p(f'{i}  {s}。')
    p('包含数据：'+data+'。')
    p('对应数据表：'+tables+'。')
    p('操作与控制：'+flow)
    p('建设安排：'+phase+'。')

h('3 从录入到结账的完整业务例子')
h('3.1 一张1000美元发票分两次付款',2)
table(['步骤','系统记录','金额变化'],[
('1 采购申请','创建业务单及订单行 船管费用中心 订单1000','只是承诺 不增加银行支出'),
('2 验收和到票','记录验收1000 发票1000 业务行确认费用','应付本金1000；正式记账示例为借费用1000 贷应付1000'),
('3 第一笔付款','付款400 实际银行手续费20；核销发票400','待付本金600；银行减少420；费用本金1000 加独立手续费20'),
('4 第二笔付款','付款600 无手续费；核销余款600','待付本金0；两次银行共减少1020'),
('5 船东结算','若合同允许报销 结算本金1000及手续费20','应收1020；是否可报销由协议决定'),
('6 船东回款','实际到账1020 并核销结算单','船东待收0；供应商付款状态和船东回款状态各自完成')],[3.3,7.2,7.1])
p('这个例子按“到票确认费用”的简化假设说明流程。若按验收应计，验收时先记费用和应计负债，到票时冲销应计并转应付；同一业务不能在验收和到票两次确认费用。税额和汇兑差额应按批准政策另记，不能塞入付款本金。')
h('3.2 银行给船长补给1000美元 船长再消费100美元',2)
p('建立内部划拨单，关联银行出账1000和现金入账1000。内部划拨前后，银行加现金的合计不变。船长消费100时新增一笔现金支出，并关联同一笔费用和发票：现金减少100，经营费用增加100。导入费用表时通过业务关联识别这笔100，不再增加第二笔费用。银行划拨手续费如有20，单独记录为银行费用，不能把1020全部当船上现金收到。')
h('3.3 同一账户替另一条船支付2600美元',2)
p('付款账户所属公司和船，与业务行中的受益船分别记录。如果本船银行替 Libra XS 支付2600及手续费20，银行减少2620；本船资金视图反映现金减少，Libra XS 的业务成本或往来按合同确认。本船经营费用视图不默认计入2600。跨主体时双方需要往来凭证，不能只把成本改到另一条船。')
h('3.4 六状态应怎样使用',2)
table(['旧表状态','业务解释及新系统处理'],[
('Accrual','旧表应计状态。新系统单独记录费用确认和应计凭证，不代表已到票或已支付。'),
('Receiving','暂译验收处理中。须由业务确认；记录验收进度，不代表收到银行款。'),
('Received','暂译已验收。记录验收完成，不能解释为船东回款。'),
('Approved','付款已批准。批准金额不自动减少银行余额。'),
('Paying','支付执行中。银行未证实前不视为真实支付完成。'),
('Completed','旧表兼容状态。正式应付结清以发票本金和有效核销为准；退款后可能重新出现待付。'),
('空白或非法状态','导入异常队列。需要明确业务类型或补充认定，不直接按零费用处理。')],[3.6,14])
p('新系统保留历史状态，同时拆开采购验收状态、发票审核状态、付款状态和船东回款状态。一张发票可能“已验收、已批准、部分付款、船东未报销”，单一状态字段无法表达这些情况。')
h('3.5 人员分工',2)
table(['角色','允许的关键操作','必须隔离的操作'],[
('业务录入 采购 船上人员','录单 上传附件 验收 查看授权船舶','不能批准自己的付款 不能记账或关账'),
('财务审核','审核归属 发票 结算 凭证草稿 对账','付款执行与审批按授权分离'),
('审批人','批准付款 资金申请 更正和重开期间','审批结果保留意见和金额范围'),
('出纳','登记真实支付 收款 银行和现金流水','不能自行批准自己的支付申请'),
('财务负责人','记账 关账 批准更正 验证恢复后的金额','重开期间需说明并留审计记录'),
('系统管理员 运维','用户授权 服务维护 备份和恢复','业务金额修改走财务流程 管理员操作同样审计')],[3.4,7.8,6.4])

# 字段定义格式为 字段名|类型与必填|中文含义及规则。所有表继承公共字段。
SCHEMA=[]
def entity(module,name,title,meaning,fields,constraint=''):
    SCHEMA.append(dict(module=module,name=name,title=title,meaning=meaning,
                       fields=[s.split('|',2) for s in fields.strip().splitlines()],constraint=constraint))
entity('M01','organization','公司主体','一条记录代表一个独立核算或管理主体。', '''
code|TXT 必填|公司代码 在系统内唯一
legal_name|TXT 必填|法定名称 由登记资料确认
short_name|TXT 必填|内部显示简称
base_currency|C3 必填 FK|本位币 引用currency.code；示例USD 由财务确认
timezone|TXT 必填|业务时区 示例Asia/Shanghai
registration_no|TXT 可空|公司注册号或税务识别号
status|ENUM 必填|ACTIVE或INACTIVE 停用后不得新增业务
''','公司改名保留审计；已用主体不得删除。')
entity('M01','vessel','船舶','一条记录代表一艘船。', '''
organization_id|ID 必填 FK|管理或核算主体 指向organization
code|TXT 必填|内部唯一船舶代码 不使用可能变化的船名作主键
imo_no|TXT 可空|IMO编号 已核实后唯一
name|TXT 必填|当前正式船名
aliases|JSON 必填|历史船名和原表别名 可为空数组
owner_party_id|ID 可空 FK|船东 指向party；启用结算前必须确认
management_start|DATE 可空|管理开始日 用于管理天数
management_end|DATE 可空|管理结束日 不早于开始日
status|ENUM 必填|ACTIVE INACTIVE等
''','原表PIONEER LAY和PIONEER LADY先列别名候选，经确认后合并。')
entity('M01','party','往来方','供应商 船东 管理公司 员工及关联方使用同一主体档案。', '''
code|TXT 必填|系统唯一往来方代码
legal_name|TXT 必填|正式名称
roles|JSON 必填|SUPPLIER OWNER EMPLOYEE AGENT RELATED等角色集合
aliases|JSON 必填|大小写 空格及历史名称变体 供导入匹配
contact_name|TXT 可空|业务联系人
contact_details|JSON 可空|电话和邮箱 按角色控制可见
registration_no|TXT 可空|注册或税务标识
payment_terms_days|INT 可空|约定账期 非负整数；没有约定时不得默认零天
status|ENUM 必填|ACTIVE或INACTIVE
''','银行收款账户信息若需要网银付款制单，应扩展独立受控档案并双人复核；首期不直接发起银行支付。')
entity('M01','currency','币种','支持原币和本位币分账。', '''
code|C3 必填 唯一|ISO币种码 示例USD CNY HKD EUR
name|TXT 必填|中文名称
minor_units|INT 必填|记账小数位 示例USD为2；按币种配置
enabled|BOOL 必填|是否允许新增交易
''')
entity('M01','fx_rate','汇率','存一单位原币折合多少目标币。', '''
from_currency|C3 必填 FK|原币 currency.code
to_currency|C3 必填 FK|目标币 currency.code
rate_date|DATE 必填|适用日期
rate|RATE 必填|大于0；目标币金额=原币金额×rate
source|TXT 必填|银行 成交价 财务批准来源等
effective_at|TS 必填|本版本生效时间
approved_by|ID 可空 FK|批准人 app_user；用于记账前必填
status|ENUM 必填|DRAFT APPROVED RETIRED
''','同币种汇率为1；已用于已记账业务的汇率不能被覆盖，新增版本。')
entity('M01','cost_center','费用中心','按组织配置五类资金用途及费用分类。', '''
organization_id|ID 必填 FK|所属主体 organization
code|TXT 必填|SHIPMANAGEMENT OWNER TAKEOVER FLAG_CHANGE DRYDOCK
name|TXT 必填|船管 船东 接船 换旗 坞修
description|TXT 可空|归类边界及允许的业务
enabled|BOOL 必填|是否启用；坞修即使为零也保留
''','同一主体内代码唯一；费用中心和费用承担方是不同字段。')
entity('M01','account_code','业务科目字典','承接原表160条船管编码和32条船员细项编码。', '''
namespace|TXT 必填|ship_management或crew_detail
code|TXT 必填|原编码 保持文本及前导零
name|TXT 必填|科目名称
name_en|TXT 可空|英文说明
parent_id|ID 可空 FK|上级account_code 不允许循环
responsible_department|TXT 可空|责任部门或PIC
gl_account_id|ID 可空 FK|正式会计科目gl_account；记账时按性质再判断
enabled|BOOL 必填|停用保留历史
''','namespace加code唯一。业务科目不是正式会计科目的替代品。')
entity('M01','gl_account','正式会计科目','在财务确认账套后启用。', '''
organization_id|ID 必填 FK|所属账套 organization
code|TXT 必填|会计科目代码 主体内唯一
name|TXT 必填|会计科目名称
parent_id|ID 可空 FK|上级gl_account
account_type|ENUM 必填|ASSET LIABILITY EQUITY INCOME EXPENSE
normal_side|ENUM 必填|DEBIT或CREDIT 正常余额方向
is_postable|BOOL 必填|是否允许直接记账 章节科目为否
requires_party|BOOL 必填|是否必须记录往来方
requires_vessel|BOOL 必填|是否必须记录船舶辅助核算
enabled|BOOL 必填|是否允许新凭证引用
''')
entity('M12','app_user','系统用户','使用受控账号登录，不直接共享财务账号。', '''
login_name|TXT 必填 唯一|登录名
display_name|TXT 必填|显示姓名
password_hash|TXT 必填|使用框架强哈希 存哈希而非明文
mfa_secret_ciphertext|TXT 可空|管理和高权限用户的加密MFA密钥
status|ENUM 必填|ACTIVE LOCKED DISABLED
last_login_at|TS 可空|最近成功登录时间
failed_login_count|INT 必填|连续失败次数 触发限流和锁定
''')
entity('M12','role','角色','权限的集合，如录入 财务审核 出纳 运维。', '''
code|TXT 必填 唯一|稳定角色码
name|TXT 必填|显示名称
description|TXT 可空|职责边界及冲突角色说明
enabled|BOOL 必填|是否可授予用户
''')
entity('M12','permission','权限点','代码发布时维护的操作目录。', '''
code|TXT 必填 唯一|例如payment.approve period.close backup.restore
name|TXT 必填|中文操作名
resource_type|TXT 必填|受保护业务对象类型
action|TXT 必填|READ CREATE APPROVE POST EXPORT等
''')
entity('M12','user_role','用户角色与数据范围','一条授权覆盖一个主体及可选一艘船。', '''
user_id|ID 必填 FK|app_user
role_id|ID 必填 FK|role
organization_id|ID 必填 FK|organization；运维跨主体需逐项授权
vessel_id|ID 可空 FK|vessel；空表示该主体所有授权业务船
valid_from|TS 必填|授权开始时间
valid_to|TS 可空|到期时间；离职立即失效
granted_by|ID 必填 FK|授权执行人app_user
''','后端每次查询和写入校验组织及船舶范围；隐藏前端按钮不足以授权。')
entity('M12','role_permission','角色权限映射','角色与操作权限多对多。', '''
role_id|ID 必填 FK|role
permission_id|ID 必填 FK|permission
''','role_id加permission_id唯一；审批自己创建的付款由服务端额外禁止。')
entity('M02','business_event','业务单','一条可追溯的经济事项，统一连接订单 发票 收付和凭证。', '''
organization_id|ID 必填 FK|承担记录责任的organization
event_no|TXT 必填|主体内唯一业务单号
vessel_id|ID 必填 FK|主要业务船vessel；跨船详见业务行
event_type|ENUM 必填|EXPENSE FUNDING DEPOSIT LOAN REFUND TRANSFER ADVANCE CREDIT等
posting_date|DATE 必填|入账所属业务日期 受期间控制
party_id|ID 可空 FK|主要往来方party；正式审核按类型要求补齐
description|TXT 必填|完整业务摘要
recognition_policy|ENUM 必填|INVOICE ACCEPTANCE或BALANCE_ONLY；费用确认触发点
document_status|ENUM 必填|DRAFT SUBMITTED APPROVED POSTED REVERSED
legacy_status|TXT 可空|原表状态 原样保留而非驱动正式核销
reversal_of_id|ID 可空 FK|冲回的原business_event
''','过账后只允许冲销或更正；同一业务不得因来自两个工作表而重复建单。')
entity('M02','event_line','业务归属明细','一条金额按受益船 费用中心和性质归属；多中心则拆多行。', '''
event_id|ID 必填 FK|business_event
line_no|INT 必填|同一业务单内唯一行号
beneficiary_vessel_id|ID 必填 FK|实际受益船vessel
cost_center_id|ID 必填 FK|cost_center
account_code_id|ID 可空 FK|account_code；审核前需补齐适用科目
responsible_party_id|ID 可空 FK|费用承担方party；结算前必填
cost_nature|ENUM 必填|OPERATING DEPOSIT LOAN INTERCOMPANY TRANSFER等
currency_code|C3 必填 FK|原币currency.code
original_amount|DEC 必填|正数或经批准的有符号更正金额
fx_rate_id|ID 必填 FK|fx_rate；币种和目标本位币须一致
base_amount|DEC 必填 计算|F01 原币金额×锁定汇率 按本位币舍入
recognition_date|DATE 可空|费用确认日；非费用事项可空
recognition_status|ENUM 必填|UNRECOGNIZED RECOGNIZED REVERSED
description|TXT 必填|本归属行的摘要
''','业务行是成本归属的唯一来源；订单 验收和发票是不同阶段，不将三个阶段金额相加作为成本。')
entity('M02','purchase_order_line','采购订单行','可选采购流程，业务单可无订单。', '''
event_line_id|ID 必填 FK|归属event_line；跨中心订单拆分对应行
requisition_no|TXT 可空|原表RQA 请购或申请编号
po_no|TXT 必填|订单号 不单独假定全系统唯一
po_line_no|TXT 必填|订单行号
order_date|DATE 必填|下单日
currency_code|C3 必填 FK|订单原币
original_amount|DEC 必填|此订单归属行的原币承诺额
fx_rate_id|ID 必填 FK|订单折算汇率
base_amount|DEC 必填 计算|F01 订单USD或本位币额
approval_status|ENUM 必填|DRAFT SUBMITTED APPROVED REJECTED CLOSED
variance_reason|TXT 可空|与业务行 发票或验收不同的原因
''','同一订单可有多次验收和多张发票；数量管理为后续可选扩展。')
entity('M02','goods_receipt','验收记录','一条代表某订单归属行的一次部分或全部验收。', '''
event_line_id|ID 必填 FK|归属event_line
po_line_id|ID 可空 FK|purchase_order_line；无订单验收可空
receipt_no|TXT 必填|验收单号或服务确认号
receipt_date|DATE 必填|验收或服务确认日
currency_code|C3 必填 FK|验收原币
original_amount|DEC 必填|本次验收金额
fx_rate_id|ID 必填 FK|验收汇率
base_amount|DEC 必填 计算|F01 本次验收本位币额
accepted_by|ID 必填 FK|验收人app_user
status|ENUM 必填|DRAFT ACCEPTED REVERSED
variance_reason|TXT 可空|超额或阶段差异说明
''','验收累计=同一订单行有效验收金额之和，超订单金额需审批；按验收应计时到票必须转回原应计。')
entity('M03','invoice','供应商发票头','一张发票或贷项凭证，不把原表文件名当作发票编号。', '''
organization_id|ID 必填 FK|organization
supplier_id|ID 必填 FK|供应商party
invoice_no|TXT 必填|真实发票号；历史缺失生成待核对内部号并标记异常
document_type|ENUM 必填|INVOICE或CREDIT_NOTE
invoice_date|DATE 必填|发票日期
due_date|DATE 可空|到期日；账龄报表缺失时显示待补
currency_code|C3 必填 FK|发票币种
fx_rate_id|ID 必填 FK|确认应付时使用的锁定汇率
net_amount|DEC 可空|未税额；税务信息不全时保持空
tax_amount|DEC 可空|税额；未知与0不同
gross_amount|DEC 必填|票面含税总额 正数；方向由document_type确定
base_amount|DEC 必填 计算|F01 票面总额折本位币
approval_status|ENUM 必填|DRAFT SUBMITTED APPROVED REJECTED VOID
verification_status|ENUM 必填|PENDING VERIFIED DISPUTED；附件及金额核验状态
''','主体 供应商 发票号及凭证类型联合防重；gross_amount=净额+税额仅在税务数据齐全时校验。')
entity('M03','invoice_line','发票明细','将一张发票分配到业务归属行，可分别核销。', '''
invoice_id|ID 必填 FK|invoice
line_no|INT 必填|发票内唯一
event_line_id|ID 必填 FK|event_line
po_line_id|ID 可空 FK|purchase_order_line
goods_receipt_id|ID 可空 FK|goods_receipt；多次验收通过拆行关联
description|TXT 必填|商品 服务或调整说明
original_amount|DEC 必填|此行含税原币额 正数
tax_rate|RATE 可空|税率 0.13表示13%；未知为空
tax_amount|DEC 可空|此行税额 不默认可抵扣
base_amount|DEC 必填 计算|F01 使用发票头汇率折算
variance_reason|TXT 可空|运费 差价 部分到票等原因
''','行原币合计等于发票总额；本位币舍入尾差分配到最后一行并留痕。应付和已付余额为查询结果，不以可手改字段保存。')
entity('M03','credit_allocation','贷项抵扣分配','贷项与原发票本金之间的多对多抵扣。', '''
credit_line_id|ID 必填 FK|类型为CREDIT_NOTE的invoice_line
invoice_line_id|ID 必填 FK|被抵扣的普通invoice_line
allocation_date|DATE 必填|抵扣生效日
currency_code|C3 必填 FK|双方相同的结算币种
original_amount|DEC 必填|本次抵扣本金 正数；冲回通过状态及关联处理
base_amount|DEC 必填|对应原应付账面值；汇兑差异单独入凭证
status|ENUM 必填|CONFIRMED REVERSED
reversal_of_id|ID 可空 FK|原credit_allocation
''','F06 累计抵扣不超过贷项可用额或发票未结清额；跨币种抵扣首期禁止，先兑换或人工批准分录。')
entity('M04','payment','付款及退款单','记录实际支付事件；现金和银行通过关联流水核实。', '''
organization_id|ID 必填 FK|organization
payment_no|TXT 必填|主体内唯一付款编号
event_id|ID 必填 FK|business_event；支持非发票押金 借款等
payee_party_id|ID 必填 FK|收款方party
payment_type|ENUM 必填|SUPPLIER DEPOSIT LOAN REFUND OTHER
payment_date|DATE 必填|实际支付或退款日期
currency_code|C3 必填 FK|实际结算币种
signed_principal|DEC 必填|出款本金为正 收到原付款退款为负
signed_fee|DEC 必填|实际付费为正 退费为负
fx_rate_id|ID 必填 FK|实际结算汇率
signed_base_principal|DEC 必填 计算|F01 结算本金折本位币 用于现金分析
signed_base_fee|DEC 必填 计算|F01 手续费折本位币；不得再计一次
bank_transaction_id|ID 条件 FK|bank_transaction 与cash_movement_id二选一；执行前可空
cash_movement_id|ID 条件 FK|cash_movement；实际完成时必须指定一个资金来源
status|ENUM 必填|DRAFT APPROVED EXECUTING SETTLED FAILED REVERSED
refund_of_id|ID 可空 FK|退款关联原payment；退款时必填
''','付款不自动改变发票费用金额；SETTLED要求真实资金记录且核销通过。已支付本金不能因付款申请审批而提前增加。')
entity('M04','payment_allocation','付款核销明细','连接付款和发票，解决一次付多票和一票分次付。', '''
payment_id|ID 必填 FK|payment
invoice_line_id|ID 必填 FK|invoice_line
allocation_date|DATE 必填|核销生效日
settlement_currency|C3 必填 FK|发票币种currency.code
signed_invoice_amount|DEC 必填|发票币种下核销本金 正常付为正 退款解除核销为负
signed_payment_amount|DEC 必填|付款币种下使用本金；同币种时等于发票核销额
carrying_base_amount|DEC 必填|本次核销原应付账面值 有符号
settlement_base_amount|DEC 必填|本次付款本金折本位币 有符号
fx_difference|DEC 必填 计算|F09 结算本位币额−核销账面值；正数通常为损失
fee_base_allocation|DEC 必填|分摊手续费 可0；仅分配已有手续费 不另生成现金
status|ENUM 必填|CONFIRMED REVERSED
reversal_of_id|ID 可空 FK|冲回的payment_allocation
''','F06 核销本金不超发票待付及付款可用本金；两边分别按各自币种校验。退款解除核销必须引用原分配。')
entity('M05','receipt','实际收款单','仅记录实际到账，借款退回与拨款按性质分开。', '''
organization_id|ID 必填 FK|organization
receipt_no|TXT 必填|主体内唯一收款单号
payer_party_id|ID 必填 FK|付款方party
receipt_type|ENUM 必填|OWNER_FUNDING REIMBURSEMENT LOAN_RETURN REFUND OTHER
receipt_date|DATE 必填|到账确认日
currency_code|C3 必填 FK|到账币种
gross_amount|DEC 必填|毛额 正数；撤销通过业务冲回
signed_fee|DEC 必填|扣费为正 退费为负
net_amount|DEC 必填 计算|F02 毛额−有符号手续费
fx_rate_id|ID 必填 FK|到账折算汇率
net_base_amount|DEC 必填 计算|F01 净到款折本位币
bank_transaction_id|ID 条件 FK|bank_transaction；现金收款时用下一个字段
cash_movement_id|ID 条件 FK|cash_movement；确认收款时资金来源二选一
status|ENUM 必填|DRAFT CONFIRMED REVERSED
event_id|ID 必填 FK|对应business_event
''','退款如已记录为payment的负付款，不再建receipt增加一次资金；按事件类型选择一种业务路径并共享银行流水。')
entity('M05','receipt_allocation','收款用途分配','一笔到账分配到多船或五类资金用途，可关联资金申请。', '''
receipt_id|ID 必填 FK|receipt
vessel_id|ID 必填 FK|资金归属船vessel
cost_center_id|ID 必填 FK|资金用途cost_center
funding_request_id|ID 可空 FK|funding_request；一笔分配可只关联一张申请 多申请则拆行
allocation_date|DATE 必填|用途分配生效日
original_amount|DEC 必填|receipt币种下净额分配额 正数
base_amount|DEC 必填 计算|F01 按收款汇率折算
fund_nature|ENUM 必填|CUSTODY MANAGEMENT_FEE LOAN REIMBURSEMENT等财务确认性质
status|ENUM 必填|CONFIRMED REVERSED
reversal_of_id|ID 可空 FK|原receipt_allocation
''','F03 用途分配不超净到款；已分配资金是管理资金池指标，不自动形成营业收入。')
entity('M09','funding_request','资金计划及补款申请','没有到账之前只形成计划。', '''
organization_id|ID 必填 FK|organization
request_no|TXT 必填|主体内唯一申请号
vessel_id|ID 必填 FK|vessel
cost_center_id|ID 必填 FK|cost_center
request_type|ENUM 必填|DAILY_PLAN CASH_SHORTFALL MANUAL
period_start|DATE 必填|计划开始日
period_end|DATE 必填|计划结束日
daily_rate|DEC 可空|日费率 DAILY_PLAN时必填
billable_days|INT 可空|计费天数 按批准的首尾日规则
currency_code|C3 必填 FK|申请币种
requested_amount|DEC 必填|F04 日费率×天数 或经审核手工缺口额
approved_amount|DEC 可空|批准金额；未批准为空不等于0
reason|TXT 必填|用途 缺口解释及人工调整原因
status|ENUM 必填|DRAFT SUBMITTED APPROVED PART_RECEIVED RECEIVED CANCELLED
''','申请实收来自receipt_allocation，批准金额−关联实收为未到账金额；关联收款须同币种，跨币种按另行批准规则处理。')
entity('M06','bank_account','银行账户','一个账户一个币种，保存受控期初。', '''
organization_id|ID 必填 FK|账户持有主体organization
account_name|TXT 必填|户名
bank_name|TXT 必填|开户行
account_number_ciphertext|TXT 必填|完整账号加密；展示按权限掩码
currency_code|C3 必填 FK|账户币种
opening_date|DATE 必填|期初余额生效日
opening_balance|DEC 必填|当日期初余额；之后只累计该日及以后流水
status|ENUM 必填|ACTIVE CLOSED
''','银行账户可能服务多船，不将一整个账户余额复制归给每条船。')
entity('M06','bank_transaction','银行流水','来自对账单的唯一资金事实。', '''
account_id|ID 必填 FK|bank_account
external_reference|TXT 可空|银行流水号 优先用于去重
booking_date|DATE 必填|银行记账日
value_date|DATE 可空|价值日 与记账日分开
sequence_key|TXT 必填|银行原序号及稳定排序键 不用导入行号随意替代
currency_code|C3 必填 FK|必须等于账户币种
signed_principal|DEC 必填|入账为正 出账为负；可0以记录独立手续费
signed_fee_movement|DEC 必填|实际扣费为负 退费为正
signed_net_movement|DEC 必填 计算|F07 有符号本金+有符号手续费变动
statement_balance|DEC 可空|银行提供的该笔后余额 不强制由系统覆盖
counterparty_text|TXT 可空|银行原文对方名称
memo|TXT 可空|原始摘要
fingerprint|TXT 必填|账户 日期 金额 参考号等规范化哈希 候选防重
source_file_id|ID 可空 FK|导入source_file
status|ENUM 必填|IMPORTED CONFIRMED DISPUTED VOID
''','fingerprint命中先阻断人工复核，不能自动删除同日同额真实付款；confirmed流水才进入正式余额，void仍留历史。')
entity('M06','bank_match','银行业务匹配','一条流水可以匹配多个业务构成，候选需人工确认。', '''
bank_transaction_id|ID 必填 FK|bank_transaction
target_type|ENUM 必填|PAYMENT RECEIPT TRANSFER JOURNAL_FEE INTEREST等
target_id|ID 必填|对应目标记录；后端验证目标真实存在且主体相符
component_type|ENUM 必填|PRINCIPAL或FEE
signed_bank_amount|DEC 必填|本次解释的账户币种有符号金额
match_method|ENUM 必填|MANUAL RULE CANDIDATE
confidence|RATE 可空|建议分值0到1 不是核销证据
confirmed_by|ID 可空 FK|确认人app_user
status|ENUM 必填|SUGGESTED CONFIRMED REVERSED
''','F08 本金与手续费分别核对，不允许正负匹配互相抵销掩盖超额；匹配只解释流水 不新增余额。')
entity('M07','cash_account','船长现金账户','每条船每种现金币种一套余额。', '''
organization_id|ID 必填 FK|organization
vessel_id|ID 必填 FK|vessel
code|TXT 必填|主体内唯一现金账户代码
currency_code|C3 必填 FK|现金币种
custodian_party_id|ID 必填 FK|保管人party
opening_date|DATE 必填|现金期初生效日
opening_balance|DEC 必填|经交接或盘点确认的期初余额
status|ENUM 必填|ACTIVE CLOSED
''')
entity('M07','cash_movement','现金流水','每笔现金收支只影响一个现金账户一次。', '''
account_id|ID 必填 FK|cash_account
movement_no|TXT 必填|现金账户内唯一流水号
movement_date|DATE 必填|现金实际收付日
movement_type|ENUM 必填|TOPUP EXPENSE RETURN ADJUSTMENT RECEIPT
signed_amount|DEC 必填|入金为正 支出为负
event_id|ID 必填 FK|business_event
internal_transfer_id|ID 可空 FK|internal_transfer；补给或退回时必填
counterparty_id|ID 可空 FK|party
counted_balance|DEC 可空|本笔后实盘余额 盘点时填写
adjustment_reason|TXT 可空|调整或盘点差异原因
status|ENUM 必填|DRAFT CONFIRMED REVERSED
''','F10 现金期末=期初+有效有符号现金流水；现金消费经payment连接发票核销，费用来源仍为业务行。')
entity('M07','internal_transfer','内部划拨','连接银行和现金的两端，首期仅支持同主体同币种。', '''
organization_id|ID 必填 FK|organization
transfer_no|TXT 必填|主体内唯一划拨编号
transfer_date|DATE 必填|划拨实际生效日
from_account_type|ENUM 必填|BANK或CASH
from_account_id|ID 必填|对应bank_account或cash_account
to_account_type|ENUM 必填|BANK或CASH
to_account_id|ID 必填|对应账户 与来源不同
currency_code|C3 必填 FK|两端相同币种
principal_amount|DEC 必填|转出及转入本金 正数
fee_amount|DEC 必填|手续费 正数或0；资金端单独扣除
status|ENUM 必填|DRAFT IN_TRANSIT CONFIRMED CANCELLED
''','F11 已确认划拨转出本金+转入本金=0；一端未到时标记在途并跟踪，不虚增收款。跨主体按往来，跨币种按兑换，首期不套此简化模型。')
entity('M08','owner_settlement','船东或关联方结算单','一张向责任主体申请报销或收回代付的结算单。', '''
organization_id|ID 必填 FK|发起结算主体organization
settlement_no|TXT 必填|主体内唯一结算号
vessel_id|ID 必填 FK|结算船vessel
debtor_party_id|ID 必填 FK|应付款给公司的船东或关联方party
period_start|DATE 必填|结算开始日
period_end|DATE 必填|结算结束日
issue_date|DATE 必填|出单日
due_date|DATE 可空|约定回款日
currency_code|C3 必填 FK|结算币种
status|ENUM 必填|DRAFT APPROVED ISSUED PART_RECEIVED SETTLED DISPUTED VOID
''','结算单不是供应商发票；单据总额和待收余额由有效明细及回款计算。')
entity('M08','owner_settlement_line','结算明细','将费用或跨船代付按责任协议计入应收。', '''
settlement_id|ID 必填 FK|owner_settlement
event_line_id|ID 必填 FK|被结算的event_line
invoice_line_id|ID 可空 FK|关联invoice_line；没有发票的批准代付可空
actual_payer_party_id|ID 必填 FK|实际出款人party
settlement_basis|ENUM 必填|COMPANY_ADVANCE OWNER_DIRECT INTERVESSEL OTHER
eligible_amount|DEC 必填|按合同可报销额 结算币种 正数
claim_amount|DEC 必填|本单申报额 不超剩余可结算额
approved_adjustment|DEC 必填|批准调减为负 调增为正 默认0
explanation|TXT 可空|差异 船东直接付或跨船说明
''','F12 可收=claim_amount+approved_adjustment；船东直接支付通常claim_amount为0，保留发生额用于说明；同一业务不得被重复结算。')
entity('M08','owner_receipt_allocation','船东回款核销','用途分配不等于报销核销，本表将真实收款核到结算行。', '''
receipt_allocation_id|ID 必填 FK|receipt_allocation；对应真实收款资金
settlement_line_id|ID 必填 FK|owner_settlement_line
allocation_date|DATE 必填|报销核销日
receipt_currency_amount|DEC 必填|占用收款币种金额 正数
settlement_currency_amount|DEC 必填|结算币种下冲减应收金额 正数
carrying_base_amount|DEC 必填|应收被核销账面本位币值
received_base_amount|DEC 必填|该部分实际收款本位币额
fx_difference|DEC 必填 计算|实际收款本位币−账面值 正数通常为收益
status|ENUM 必填|CONFIRMED REVERSED
reversal_of_id|ID 可空 FK|原owner_receipt_allocation
''','核销不超结算待收或该收款用途分配未核销额；同币种金额相等；两边分别校验，不能通过重复用途分配重复报销。')
entity('M10','inventory_month','金额型库存月结','一条记录对应一船 一月 一种库存和一种币种。', '''
organization_id|ID 必填 FK|organization
vessel_id|ID 必填 FK|vessel
month|DATE 必填|月份第一天
inventory_type|ENUM 必填|FOOD或BONDED
currency_code|C3 必填 FK|估值币种
opening_value|DEC 必填|期初库存估值 非负
purchase_value|DEC 必填 计算|当月有效PURCHASE变动合计
transfer_net|DEC 必填 计算|当月调入减调出
other_adjustment|DEC 必填 计算|批准的其他库存调整 有符号
closing_count_value|DEC 可空|实盘期末库存估值；未盘点保持空
consumption_value|DEC 可空 计算|F13 期初+采购+调拨净额+调整−期末
income_value|DEC 必填|本月伙食或招待品收入 经审核录入或关联业务
surplus_value|DEC 可空 计算|F14 收入−实际耗用
counted_by|ID 可空 FK|盘点app_user；结账前必填
status|ENUM 必填|OPEN COUNTED CLOSED
''','船 月 类型 币种联合唯一；结账要求非空盘点值。按金额倒轧耗用无法独立识别盘亏，需数量台账后才能分开。')
entity('M10','inventory_movement','库存金额变动','承接采购 调拨及批准调整，不等同于现金付款。', '''
inventory_month_id|ID 必填 FK|inventory_month
movement_date|DATE 必填|必须在对应月份
movement_type|ENUM 必填|PURCHASE TRANSFER_IN TRANSFER_OUT ADJUSTMENT
signed_value|DEC 必填|增加为正 减少为负
event_line_id|ID 可空 FK|相关event_line；采购确认后要求有关联
description|TXT 必填|采购或调整说明
approved_by|ID 可空 FK|调整审批app_user
status|ENUM 必填|DRAFT CONFIRMED REVERSED
''','现金付款关联同一采购业务，不再生成第二笔库存采购；本方案不包含商品数量 批次和成本计价。')
entity('M09','settlement_term','合同与计算参数版本','存已批准的汇率方向 分成和估算策略。', '''
organization_id|ID 必填 FK|organization
vessel_id|ID 可空 FK|适用vessel 空为主体通用
term_code|TXT 必填|同参数代码允许多个生效版本
effective_from|DATE 必填|生效日
effective_to|DATE 可空|失效日
term_type|ENUM 必填|SHARE ESTIMATE DAY_COUNT RECOGNITION ROUNDING
parameters|JSON 必填|参数对象；例如share_ratio=0.45 low_policy=ZERO day_mode=INCLUSIVE
formula_version|TXT 必填|服务端公式版本 不执行任意用户脚本
approved_by|ID 必填 FK|财务或业务负责人app_user
status|ENUM 必填|DRAFT APPROVED RETIRED
''','同一适用范围与代码的已批准生效区间不可重叠；原表自由常数未经确认不自动生效。')
entity('M09','estimate_item','费用估算项','保留低 中 高估算和缺失值，不自动记账。', '''
vessel_id|ID 必填 FK|vessel
term_id|ID 可空 FK|settlement_term；自动估算必须引用批准版本
description|TXT 必填|估算事项
catalogue|TXT 必填|Shipowner Flag DOC等分类
agreement_type|TXT 可空|Take-over或OPEX等协议归属
supplier_id|ID 可空 FK|party
requisition_no|TXT 可空|原RQA
po_no|TXT 可空|原订单号
currency_code|C3 必填 FK|估算原币
original_amount|DEC 可空|未提供保持空 不默认0
rate_value|RATE 必填|大于0
rate_direction|ENUM 必填|MULTIPLY或DIVIDE；原分析模板为除法
base_amount|DEC 可空 计算|F01 方向为乘时原币×rate 方向为除时原币÷rate
place_factor|DEC 可空|地区因素 仅在批准公式引用时参与
overtime_factor|DEC 可空|加班因素 仅在批准公式引用时参与
low_amount|DEC 可空|F15 按批准策略取0或估算基数 可手工有理由覆盖
mid_amount|DEC 可空|无统一公式 手工批准值 不自动取高低平均
high_amount|DEC 可空|F15 多数取估算基数 具体按版本
override_reason|TXT 可空|人工覆盖原因
status|ENUM 必填|DRAFT APPROVED CONVERTED CANCELLED
''','空值表示未知，零表示确认零；估算转正式业务需新建business_event并保留source映射。')
entity('M12','attachment','附件对象','元数据在数据库，文件放不可变对象目录。', '''
organization_id|ID 必填 FK|organization
original_name|TXT 必填|上传文件原名或旧表中出现的文件名
object_key|TXT 可空|实际文件存储键；只有缺档占位可空
content_sha256|TXT 可空|实际文件字节SHA256；补档后必填
byte_size|BIGINT 可空|文件长度 补档后必填
mime_type|TXT 可空|服务端识别类型
document_type|ENUM 必填|INVOICE BANK_SLIP APPROVAL CREDIT_NOTE OTHER
uploaded_by|ID 可空 FK|app_user；旧表占位可空
uploaded_at|TS 可空|实际上传完成时间
status|ENUM 必填|MISSING QUARANTINED AVAILABLE RETAINED
''','文件先写入并校验后才提交AVAILABLE引用；覆盖上传生成新对象键。关账附件和源文件禁止直接覆盖或删除。')
entity('M12','attachment_link','附件业务关联','一份文件可关联多笔业务，一笔业务可有多份文件。', '''
attachment_id|ID 必填 FK|attachment
entity_type|TXT 必填|已注册业务对象类型
entity_id|ID 必填|对应业务ID 后端及迁移校验存在性
purpose|ENUM 必填|INVOICE_PROOF PAYMENT_PROOF APPROVAL OTHER
''','关联记录与文件访问范围必须同时满足组织和船舶授权；多态外键由服务层及完整性扫描校验。')
entity('M12','source_file','源文件与导入批次','留存原Excel和银行对账单字节。', '''
organization_id|ID 必填 FK|organization
original_name|TXT 必填|原文件名
sha256|TXT 必填|源文件字节SHA256
object_key|TXT 必填|原文件只读存储键
file_type|ENUM 必填|WORKBOOK BANK_STATEMENT OTHER
imported_by|ID 必填 FK|app_user
imported_at|TS 必填|导入时间
parser_version|TXT 必填|提取和清洗程序版本
status|ENUM 必填|STAGED REVIEWED IMPORTED REJECTED
''','同主体同文件哈希防重复整批入账；文件存储纳入附件备份和恢复清单。')
entity('M12','source_cell_mapping','原表单元格到系统映射','保留所有源单元格，包括暂未采用的公式和空状态。', '''
source_file_id|ID 必填 FK|source_file
sheet_name|TXT 必填|工作表名
cell_address|TXT 必填|如M70或AE15
raw_type|TXT 必填|number text date formula error等
raw_value|JSON 可空|原字面值或原缓存值；可区分空文本 空值 数字零
raw_formula|TXT 可空|精确原Excel公式 不执行宏或外部内容
target_type|TXT 可空|映射目标类型；纯归档单元格可空
target_id|ID 可空|映射目标ID；与target_type同时有值
transform_rule|TXT 必填|归并 换算 保留或排除规则及版本
review_status|ENUM 必填|PENDING CONFIRMED EXCLUDED ARCHIVED
''','一个单元格可对应多个归属行；多页重复业务可映射同一目标。目标类型和ID由服务层验证。')
entity('M12','migration_issue','迁移异常与待办','逐项处理币种缺失 重复候选 状态冲突及差额。', '''
source_file_id|ID 必填 FK|source_file
source_ref|TXT 必填|工作表 行或单元格定位
issue_type|ENUM 必填|DUPLICATE CURRENCY STATUS AMOUNT FORMULA MISSING_FILE BALANCE OTHER
severity|ENUM 必填|BLOCKING WARNING INFO
description|TXT 必填|事实 影响和待补资料
resolution|TXT 可空|确认依据和处理结果 不能只填已解决
assigned_to|ID 可空 FK|负责人app_user
resolved_by|ID 可空 FK|复核人app_user
resolved_at|TS 可空|关闭时间
status|ENUM 必填|OPEN INVESTIGATING RESOLVED ACCEPTED_LEGACY
''','ACCEPTED_LEGACY只允许保留在兼容层；未解决的关键异常不得直接进入正式凭证。')
entity('M12','workflow_action','审批和状态事件','保存每次提交 审批 拒绝 撤回及核销确认。', '''
organization_id|ID 必填 FK|organization
entity_type|TXT 必填|审批对象类型
entity_id|ID 必填|审批对象ID
action|ENUM 必填|SUBMIT APPROVE REJECT WITHDRAW CONFIRM REOPEN
from_status|TXT 必填|操作前状态
to_status|TXT 必填|操作后状态
actor_id|ID 必填 FK|app_user
acted_at|TS 必填|操作时间
comment|TXT 可空|意见 拒绝和重开时必填
approved_limit|DEC 可空|本次批准金额及原币范围
request_id|TXT 必填|关联请求及幂等键
''','追加写入 不覆盖旧审批；审批和核心状态变更在同一数据库事务内完成。')
entity('M12','audit_log','审计日志','记录查询导出之外的关键变更和敏感操作。', '''
organization_id|ID 可空 FK|未登录事件可空 已登录业务事件必须有主体
actor_id|ID 可空 FK|app_user；系统任务单独标记
actor_type|ENUM 必填|USER SYSTEM SERVICE
occurred_at|TS 必填|UTC操作时间
action|TXT 必填|登录 修改 授权 导出 下载 记账 恢复等
entity_type|TXT 可空|对象类型
entity_id|ID 可空|对象ID
before_data|JSON 可空|操作前必要字段 敏感数据脱敏 不记录密码与密钥
after_data|JSON 可空|操作后必要字段
client_ip|TXT 可空|来源IP
request_id|TXT 必填|请求追踪号
result|ENUM 必填|SUCCESS DENIED FAILED
''','业务用户不可删改；日志异步复制到独立仓库。数据库管理员仍有技术权限，外部副本和权限分离提供额外证据。')
entity('M11','accounting_period','会计及管理期间','期间锁定按主体执行，船舶报表在同一期间内过滤。', '''
organization_id|ID 必填 FK|organization
period_code|TXT 必填|例如2026-10 主体内唯一
start_date|DATE 必填|开始日 含当日
end_date|DATE 必填|结束日 含当日
status|ENUM 必填|OPEN CLOSING CLOSED
closed_by|ID 可空 FK|关账app_user
closed_at|TS 可空|关账时间
snapshot_batch|TXT 可空|本次关账报表批次
reopen_reason|TXT 可空|重开需审批及说明
''','同主体期间不重叠；CLOSED期间禁止新增及修改已确认财务影响，后期更正或批准重开。')
entity('M11','journal_entry','会计凭证头','经会计政策确认后由业务生成草稿，审核再记账。', '''
organization_id|ID 必填 FK|organization
entry_no|TXT 必填|主体 期间内唯一凭证号
period_id|ID 必填 FK|accounting_period
posting_date|DATE 必填|凭证日期 必须在所属期间
event_id|ID 可空 FK|business_event；期初及人工调整可空但须说明
source_stage|TXT 必填|ACCRUAL INVOICE PAYMENT RECEIPT OPENING ADJUSTMENT等
generation_key|TXT 必填 唯一|业务阶段及版本幂等键 防止同业务重复制证
description|TXT 必填|凭证摘要
status|ENUM 必填|DRAFT REVIEWED POSTED REVERSED
reviewed_by|ID 可空 FK|审核app_user
posted_by|ID 可空 FK|记账app_user
posted_at|TS 可空|记账时间
reversal_of_id|ID 可空 FK|冲销journal_entry
''','F20 借贷相等才可记账；历史兼容差额不得生成没有依据的调平凭证。')
entity('M11','journal_line','会计凭证明细','正式总账余额只由已记账凭证明细形成。', '''
entry_id|ID 必填 FK|journal_entry
line_no|INT 必填|凭证内唯一行号
gl_account_id|ID 必填 FK|可记账gl_account
vessel_id|ID 可空 FK|船舶辅助核算vessel；科目要求时必填
party_id|ID 可空 FK|往来方party；往来科目要求时必填
cost_center_id|ID 可空 FK|费用中心cost_center
currency_code|C3 必填 FK|交易原币
original_debit|DEC 必填|原币借方 非负
original_credit|DEC 必填|原币贷方 非负
fx_rate_id|ID 必填 FK|锁定汇率 同币种为1
base_debit|DEC 必填|本位币借方 非负
base_credit|DEC 必填|本位币贷方 非负
description|TXT 必填|行摘要及业务依据
''','每行借贷只允许一侧大于0；整凭证本位币借方合计等于贷方。期初也使用OPENING凭证，不另藏余额表。')
entity('M11','report_snapshot','报表快照','冻结指定范围和口径的输出，支持日后复现。', '''
organization_id|ID 必填 FK|organization
vessel_id|ID 可空 FK|vessel 空为主体汇总
report_code|TXT 必填|FUNDS COST AP_AR LEGACY_SOFP GL等
period_start|DATE 可空|期间型报表开始日
as_of_date|DATE 必填|显式截止日 禁止以TODAY改动历史
currency_code|C3 必填 FK|输出币种
rule_version|TXT 必填|公式及业务规则版本
data_cutoff_at|TS 必填|生成时纳入的交易提交边界
snapshot_version|INT 必填|同范围报表修订版次
generated_by|ID 必填 FK|app_user
generated_at|TS 必填|生成时间
status|ENUM 必填|DRAFT APPROVED LOCKED SUPERSEDED
content_sha256|TXT 必填|指标规范化内容哈希
''','兼容快照和正式报表分report_code；旧版冻结，更正发布新版本。')
entity('M11','report_metric','报表指标明细','保存每张快照中的指标和可追溯计算输入。', '''
snapshot_id|ID 必填 FK|report_snapshot
metric_code|TXT 必填|FUNDING_TOTAL BANK_BALANCE等稳定指标码
cost_center_id|ID 可空 FK|费用中心维度
value|RAW 可空|原精度数值 正式报表按已批准舍入输出；未知为NULL
unit|TXT 必填|USD CNY DAYS等
formula_code|TXT 必填|对应本方案F编号或历史R编号
inputs|JSON 必填|使用的聚合范围 原始来源ID或归档清单键
quality_status|ENUM 必填|VERIFIED UNRECONCILED MISSING LEGACY
''','汇总字段是可重建的快照，不作为可手工修改的实时账本。')
entity('M13','backup_run','备份运行记录','备份目录在独立仓库，数据库仅保存一份查询副本。', '''
job_code|TXT 必填|任务名称 DB_FULL DB_DIFF FILES OFFSITE等
repository_id|TXT 必填|目标仓库稳定标识 不能包含密码
backup_type|ENUM 必填|FULL DIFFERENTIAL WAL FILES CONFIG OFFSITE
started_at|TS 必填|开始时间
completed_at|TS 可空|完成时间
status|ENUM 必填|RUNNING SUCCESS FAILED VERIFIED
manifest_key|TXT 可空|独立仓库中的备份清单和工具原生日志
recoverable_until|TS 可空|校验连续WAL及附件覆盖后的可恢复截止时间
byte_count|BIGINT 可空|本次写入字节数
checksum_status|ENUM 必填|PENDING PASS FAIL
expires_at|TS 可空|按保留策略计算 不独立删除备份链
error_summary|TXT 可空|失败原因和告警状态
''','SUCCESS只表示任务执行完成，VERIFIED要求校验；能否恢复以restore_drill实测为准。')
entity('M13','restore_drill','恢复演练和事故记录','保存从备份恢复到可用业务的证据。', '''
drill_no|TXT 必填|唯一演练编号
scenario|ENUM 必填|HOST_LOSS PITR FILE_LOSS SITE_LOSS RANSOMWARE
backup_manifest_key|TXT 必填|使用的独立备份清单
target_recovery_at|TS 必填|选择恢复到的时间点
incident_started_at|TS 必填|事故或演练开始时间
service_ready_at|TS 可空|业务验证通过并恢复服务时间
last_recovered_transaction_at|TS 可空|已恢复最后一笔确认业务时间
rpo_seconds|INT 可空 计算|相对事故时刻的实际数据回退时长
rto_seconds|INT 可空 计算|F25 服务可用时刻−事故开始时刻
validation_result|JSON 必填|余额 核销 凭证 附件抽查 哈希 登录和权限验证
executed_by|ID 必填 FK|运维app_user
accepted_by|ID 可空 FK|业务或财务验收app_user
status|ENUM 必填|PLANNED RUNNING PASSED FAILED
lessons|TXT 可空|差距 修复动作与责任人
''','目标RPO由独立监测的归档滞后和最后可恢复业务共同判断，不能以备份任务启动时间替代。')

for e in SCHEMA:
    if e['name'] in ['receipt','receipt_allocation','credit_allocation','owner_receipt_allocation']:
        e['fields'].append(['effect_sign','INT 必填','正常记录为1 冲回记录为-1；统计金额=金额×effect_sign，冲回必须关联原记录'])
        if e['name']=='receipt':
            e['fields'].append(['reversal_of_id','ID 可空 FK','冲回关联原receipt；本金额与费用和分配同步反向'])
    if e['name']=='invoice':
        for row in e['fields']:
            if row[0]=='due_date': row[2]+='；已批准账期时可计算为invoice_date+约定天数，手工覆盖须记录理由'

h('4 每个模块包含哪些数据及其关联')
h('4.1 数据表设计的三个层次',2)
p('第一层为主数据：公司、船、往来方、科目和汇率，解决“是谁、属于哪里、按什么标准”。第二层为业务事实：费用、发票、收付款、账户流水和核销，解决“发生了什么”。第三层为控制与结果：审批、审计、关账、快照和备份，解决“谁确认、怎样复现、怎样恢复”。')
p(f'本方案定义{len(SCHEMA)}张逻辑数据表。每张表的专用字段和所有公共字段在第9章列明。它们是拟建系统的数据设计，不表示当前项目已经存在这些数据库表。角色和权限等基础表使用框架实现时可以映射为框架原生表，业务含义和约束保持一致。')
for mid,title,*_ in MODULES:
    es=[e for e in SCHEMA if e['module']==mid]
    if es:
        h(mid+' '+title+'的数据',2)
        table(['数据表','中文名称及一条记录代表什么'],[(e['name'],e['title']+'；'+e['meaning']) for e in es],[6.5,11.1])
p('上表按主责模块列一次；跨模块可以共享。例如发票归属M03，付款核销归属M04，但M03需要读取核销结果计算待付。M08也会读取M05实际收款，不能再录入一份“已收金额”冒充回款。')
h('4.2 关键关联关系',2)
table(['关联','关系及作用'],[
('公司 → 船舶 账户 期间','1对多；每笔跨表引用同时校验主体一致性'),
('业务单 → 业务行 → 订单行及验收','1对多；业务行统一成本性质与受益船，订单和验收保留各自阶段金额'),
('发票 → 发票行 → 业务行','1对多及多对1；一笔业务可分次到票，一张发票可覆盖多个业务'),
('付款 ↔ 发票行','通过payment_allocation形成多对多；保存本次核销本金及退款冲回'),
('银行流水 ↔ 收付款及划拨','通过bank_match逐笔解释；同一流水可拆本金和手续费'),
('实际收款 → 资金用途分配','1对多；分到船和费用中心，按资金性质分类'),
('资金用途分配 ↔ 船东结算行','通过owner_receipt_allocation多对多；实际到账才能冲减待收'),
('内部划拨 → 银行或现金两端','两端本金相互对应；不增加经营费用 手续费另计'),
('业务事项 → 会计凭证 → 凭证行','同一业务可在应计 到票 支付等阶段生成不同凭证；每阶段幂等'),
('源单元格 ↔ 正式业务','映射表允许多个原表坐标对应一个目标；每条归并有理由'),
('报表快照 → 指标明细','1对多；快照固定截止日 数据边界与规则版本')],[6.4,11.2])
h('4.3 哪些数据不应重复保存为可编辑金额',2)
p('待付款、待收款、银行计算余额、现金计算余额、费用中心净额和凭证余额，应根据有效事实计算，页面上不能任意修改。报表快照可固化计算结果用于审计，但不能反过来修改交易。订单、验收和发票金额则必须分别保存，因为它们回答不同问题，可能存在合法差异。')
p('“受益船、费用中心、责任方、业务性质”从业务行统一引用；发票和付款不各自复制一套可自由修改的成本归属。正式成本取达到确认条件的业务金额及已批准调整，三阶段差异会阻止正式记账或要求差异说明。')

FORMULAS=[
('F01','原币折算','base=ROUND(original×rate,目标币小数位)。除法方向先规范化rate=1÷原除数；历史估算亦可保留原方向显示。','1000 CNY×0.149=149 USD。记录汇率日 来源 方向和版本；本例仅示意。'),
('F02','到账净额','net=gross−signed_fee；signed_fee>0为扣费，<0为退费。','毛额1000 扣费20 净到款980；退费20则signed_fee=-20。银行流水另用有符号变动规则。'),
('F03','资金用途分配','unallocated=receipt.net−Σ有效用途分配；全部分配时Σ分配=净到款。','980分船管800 接船180。允许暂未分配，但余额必须显式展示且不可小于0。'),
('F04','计划拨款与到账','requested=daily_rate×billable_days；unreceived=approved−Σ关联有效实收分配。','5000×31=155000；到账前银行及实收不变；计划超收另列差异。'),
('F05','阶段差异','验收差额=Σ有效验收−订单额；发票差额=Σ有效到票−Σ有效验收。','必须在相同币种或锁定的比较汇率下比较；差异不是自动错误。'),
('F06','发票待付本金','outstanding=有效发票本金−有效贷项抵扣−Σ有符号付款核销本金。','1000−0−400=600；再付600后为0；退回400并解除核销后恢复400。超付按预付或待分配，不用MAX(0)掩盖。'),
('F07','银行变动及余额','signed_net=signed_principal+signed_fee_movement；balance(T)=opening+Σ截至T的有效signed_net。','付2600费20：-2600+(-20)=-2620；原款及费全退：+9500+20=+9520。'),
('F08','银行未解释额','unmatched_principal=流水有符号本金−Σ已确认本金匹配；fee同样单独计算。','必须分别约束出入方向和金额；不同流水的正负差不能互抵。计算余额−银行对账单余额另作对账差。'),
('F09','付款汇兑差异','fx_difference=核销部分付款本位币额−原应付账面本位币额。','付款折148 原应付账面149 差-1；通常为汇兑收益1，凭证方向经财务确认。'),
('F10','船长现金余额','ROB(T)=期初现金+Σ截至T有效现金有符号变动。','原快照现金入金64815−支出34907=29908。盘点差异=实盘−计算余额，审批后单独调整。'),
('F11','内部划拨','同币种转入本金−转出本金=0；资金总变化仅为实际外部手续费。','银行-1000 现金+1000 无费时总资金不变；手续费20则总资金减少20。'),
('F12','船东待收','receivable=Σ申报金额+Σ批准调整−Σ有效回款核销结算币种金额。','申报1020 回款400 待收620。船东直接付通常不再形成公司待收。'),
('F13','库存实际耗用','consumption=opening+purchases+transfer_in−transfer_out+adjustments−closing_count。','100+300−80=320。盘点值未知时耗用未知；不得把空盘点当0。'),
('F14','库存业务月结','surplus=本月收入−实际耗用；next_opening=本月期末盘点值。','收入400 耗用320 结余80。现金另结，不把库存当现金。'),
('F15','低中高估算及分成','低值按批准策略取0或base；高值按策略取base；中值人工审核；分成=批准可分成基数×比例。','45%和55%只有合同确认后启用；各比例合计=100%，尾差按约定分配。'),
('F16','费用中心资金净额','closing=opening+有效净收款用途分配−已确认归属成本±批准用途调整。','这是管理资金结余；与现金账户余额含义不同。押金 借款 划拨按性质单列，不能全部扣经营费用。'),
('F17','真实日均经营成本','daily_OPEX=期间已确认经营费用÷批准计费天数；天数为0时不计算。','含首尾日：end−start+1；不含首日：end−start。旧表371554.52÷91=4083.0167是资金日均，不能替代费用日均。'),
('F18','资金预测缺口','shortfall=max(0,预测期内应付现金+已批准其他支出−可用银行−可用现金−已确认预期到账)。','预期到账仅用于预测，不入实收；指定预测窗口 币种和受限资金，避免把全部应付都视作今天到期。'),
('F19','应付及应收账龄','overdue_days=max(0,as_of_date−due_date)；余额按0天 1至30 31至60 61至90 90天以上分组。','缺到期日显示待补，不自动归为未逾期；历史截止日只扣该日前有效核销。'),
('F20','凭证平衡与科目余额','每凭证Σbase_debit=Σbase_credit；借方余额=期初借贷差+本期借方−本期贷方。','贷方正常科目反向展示。正式凭证差额必须为0本位币最小单位，不能容忍1美元差额。'),
('F21','六状态历史分桶','普通费用bucket(s)=IF(legacy_status=s,invoice_usd,0)。CTM应计取PO USD Receiving取GRN 其余取发票。','仅兼容层使用；原表空文本另保留。正式部分付款余额使用F06。'),
('F22','旧表总费用与资金差','legacy_cost=Σ六状态金额+全部历史手续费；legacy_net=371554.52−398259.751018376。','旧净额=-26705.231018376；显示-26705.23。不能直接称为利润或法定权益。'),
('F23','旧表资金平衡与现金覆盖','difference=实收−银行−CTM−Completed含费；coverage=银行+CTM−五个未完成状态金额。','资金差=-2790.581998376；coverage=-23914.64902。兼容原口径，不用差额自动调账。'),
('F24','金额舍入与差额','新增正式金额默认十进制ROUND_HALF_UP到币种小数位；ε=0.01仅用于USD匹配提示。','有0.01差额仍须显示及按批准尾差规则处理；凭证借贷最终严格相等。历史原精度独立保存。'),
('F25','恢复实测','RTO=service_ready_at−incident_started_at；实际回退时长=事故时刻−已证实恢复业务的最后时间。','在已知演练交易连续序列下验证RPO；低业务量时结合连续WAL和附件检查点，不能只看最后交易年龄。'),
]
h('5 关键计算公式与口径')
p('所有公式在服务端统一执行。金额采用十进制定点数；舍入时点、费用确认时点、币种和截止日期随规则版本保存。除历史复现外，报表只能纳入已批准或已确认的有效记录，并排除截止日之后的影响；截至日余额必须使用截至该日的核销和冲回记录，不能套用今天的最终状态。')
for code,title,formula,notes in FORMULAS:
    h(code+' '+title,2); p('公式：'+formula); p('说明：'+notes)
h('5.1 旧表和新系统应如何同时存在',2)
p('历史兼容报表按原缓存和原公式复现，标明历史口径及异常；正式业务报表按已经批准的新口径计算。两者金额不一致时提供桥接表，分别列出跨船代付、备用金转账、船东直接付款、状态缺失、手续费及截止期调整。桥接表的每一项都必须回到业务和凭证，不以“系统差额”代替原因。')
p('原表平衡阈值有的为绝对值小于1，有的允许等于1；新系统不沿用此含混边界。匹配容差只生成提示和待办，不自动产生财务分录。正式凭证在本位币最小记账单位上必须严格借贷相等。')

h('6 技术栈拆分与选型')
p('推荐采用模块化单体应用：业务模块在同一个受控后端中分层实现，使用一个事务数据库和独立文件存储。这个形态便于局域网部署、逐笔追溯和整体恢复。首期不需要为每个模块拆成单独服务；规模扩大后可按实际瓶颈拆分。')
table(['层次','建议技术','用途及选择原因'],[
('前端','Vue 3 TypeScript Vite Element Plus','表单 台账 核销页面和报表；金额以字符串传输和显示，核心财务计算在后端'),
('后端','Python 3.13 Django 5.2 LTS Django REST Framework','业务接口 事务 权限 审批和后台管理；Django 5.2 LTS支持期至2028年4月，投产前锁定受支持补丁版'),
('主数据库','PostgreSQL 18','外键 事务 约束 行锁 定点数 JSON及WAL时间点恢复；主从大版本保持一致'),
('财务计算','Python Decimal PostgreSQL NUMERIC','避免二进制浮点误差；统一公式服务 规则版本和舍入策略'),
('文件存储','本地独立文件卷 不可变对象键','发票 回单 原表导入文件；数据库存元数据，文件通过鉴权接口访问，不直接共享给员工写入'),
('后台任务','独立worker及systemd timer','按source_file和backup_run状态调度导入与备份，固定任务版本化配置；任务幂等和失败重试；复杂队列作为后续扩展'),
('Web入口','Nginx 内部HTTPS Gunicorn','统一局域网访问 反向代理 请求大小限制；使用公司内部CA并下发信任'),
('运行环境','Ubuntu Server 24.04 LTS Docker Compose','固定镜像和持久卷；数据库也可用受管宿主机服务，选定一种后形成运维手册'),
('数据库备份','pgBackRest 基础备份 差异备份 连续WAL','支持指定时间点恢复；由systemd timer调度，备份仓库不放在生产数据库同盘'),
('附件与配置备份','restic 加密版本化仓库','增量备份 不可变文件和配置清单；独立保留周期与离线密钥'),
('监控与日志','Prometheus Grafana Alertmanager 加独立日志归档','监测磁盘 复制滞后 WAL归档 备份覆盖和服务健康；首期可简化界面，但关键告警必需'),
('测试与交付','pytest API测试 Playwright 版本锁文件 离线镜像包','验证金额 核销 权限 迁移和恢复；持续集成可用公司现有平台，测试库与生产隔离')],[2.6,5.8,9.2])
p('这些是建议选型，非当前已安装环境。前端和辅助组件在开发启动时核验官方支持及兼容性并锁版本，不在生产自动升级。数据库物理恢复要求匹配大版本及工具版本；升级前必须先验证旧备份可恢复并保留回退环境。')
h('6.1 后端应拆分的代码模块',2)
table(['模块包','实现职责'],[
('identity masterdata','用户 角色 组织 船舶 业务字典和汇率'),
('procurement expenses payables','业务单 归属 验收 发票 贷项及应付'),
('payments receipts treasury','付款 回款 核销 银行 现金 划拨及对账'),
('settlements planning inventory','船东结算 跨船往来 资金计划和库存'),
('ledger reporting periods','凭证 期间锁定 公式服务及报表快照'),
('migration documents audit operations','旧表导入 附件 审计 备份任务状态及运维接口')],[6.1,11.5])
h('6.2 必须落实的工程规则',2)
p('核销、审批、记账和期间锁定使用数据库事务。部分付款核销时锁定发票可用余额和付款剩余额，防止两个用户同时核销导致超额；API使用幂等键阻止重复提交。每张业务表增加组织过滤，所有外键在服务端检查主体、船舶范围和币种。')
p('过账和关账后的财务事实追加冲销、更正版本，保留旧记录。业务ID使用UUID，单据编号另行生成；记录UTC时间，页面按公司时区展示。日期与时刻分开，月底不能因时区自动跨期。数据库迁移先在恢复环境演练，更新失败按发布包及迁移回退方案处理。')
p('局域网运行不调用外网字体、CDN、在线汇率或云端登录。安装依赖、容器镜像、字体、备份工具和证书信任包提前下载并校验哈希，保存离线安装包及软件清单；后续补丁经测试后在维护窗口更新。飞书用于方案文档导入，不作为核心财务系统运行依赖。')

h('7 局域网部署 备份与容灾方案')
h('7.1 推荐部署拓扑',2)
p('员工电脑通过内部域名访问HTTPS入口 → A生产服务器上的Nginx和应用 → PostgreSQL主库及独立附件卷。A只向B待机服务器复制数据库和不可变附件；C备份节点保存独立备份仓库；D为断开连接、存放异地的加密介质。A、B、C不能只是同一台物理机上的三个虚拟机。')
table(['节点','建议起步资源','职责和故障边界'],[
('A 生产服务器','8核 32GB内存 系统SSD；数据库及附件用镜像数据盘','运行应用和主库；镜像盘降低单盘故障停机，但不能代替独立备份'),
('B 待机服务器','8核 32GB内存 足够容纳全量数据的SSD','独立物理主机；PostgreSQL异步流复制，应用镜像预置，附件增量复制；平时不对用户写入'),
('C 备份节点','4核 16GB内存 备份容量按公式核算 可从4TB可用容量规划','与生产分账号 分权限；保存pgBackRest及restic仓库；尽量不同机柜电源'),
('D 异地离线介质','至少3套加密移动盘或等效离线介质','一套用于复制 一套已在异地 一套轮换；不得同时全部接到生产网'),
('网络与供电','千兆以上网络 管理VLAN UPS 独立内部DNS和时间服务','数据库和备份端口仅对指定主机开放；UPS容量须支持测算后的安全关机时间')],[3,5.4,9.2])
p('资源按首期约10至30名用户、5艘以内船舶、附件逐年增长的假设给出，不是已测算容量承诺。开发前采集实际并发、日增附件、数据库规模和WAL产生量，再调整内存和存储。测试环境独立；生产服务器不同时承担开发试验。')
h('7.2 网络和账号安排',2)
p('员工网只开放HTTPS 443；数据库5432、SSH、复制和备份端口只允许应用与管理网指定主机。附件目录不做员工可写的SMB共享。使用内部域名、内部CA、强密码和高权限账号MFA；密码、数据库账号、加密密钥不放在代码仓库或明文导出文件中。')
p('应用数据库账号不使用超级用户；备份账号、复制账号和日常运维账号分开。备份节点只接受必要数据及受限工具操作，删除备份和保留策略维护使用独立管理凭据。附件仓库可采用rest-server的append-only入口，清理由隔离管理端执行；离线介质提供另一层隔离。数据库备份仓库按pgBackRest所需权限单独隔离，不能假定普通目录权限就具备不可变能力。')
h('7.3 需要一并保护的数据',2)
table(['保护对象','具体内容','遗漏后的后果'],[
('数据库','全部业务表 用户 权限 期间 凭证 审计及报表快照','金额或核销无法恢复'),
('附件与源文件','发票 回单 审批材料 原Excel 银行导入文件及对象哈希','数据库有记录但凭证打不开'),
('运行配置','Compose环境说明 Nginx 数据库配置 定时任务 防火墙 内部DNS及证书恢复说明','有数据但服务启动失败或用户无法访问'),
('程序与安装包','应用版本 镜像摘要 锁文件 数据库迁移 依赖和工具离线包','无法恢复同版本软件和数据库结构'),
('密钥与恢复凭据','仓库加密密钥 应用字段加密密钥 证书私钥 应急管理员凭据','备份存在却无法解密或账号字段无法读回'),
('清单及恢复手册','备份ID 文件清单 一致性检查点 WAL连续性 工具版本 恢复步骤 联系人','事故发生后不知道从哪份备份恢复')],[3.2,8.3,6.1])
p('密钥用单独加密封存件保存，两名授权保管人分别持有恢复所需材料；一份与异地介质分开放置。每季度实际解密验证一次。密钥副本不只放在需要该密钥解密的仓库里。权限变更和人员离职时执行凭据轮换，并验证旧备份仍可解密。')
h('7.4 数据库与附件必须恢复到同一业务边界',2)
p('数据库WAL可以把数据库恢复到时间点T，但不会自动恢复发票PDF、配置或密钥。系统因此采用不可变文件对象：先写文件、原子完成并计算哈希，再提交数据库AVAILABLE引用；新版本文件使用新键，原文件不覆盖。在最长备份保留和财务归档期内，不物理删除仍可能被历史数据库引用的对象。')
p('每5分钟生成一个联合恢复检查点：短暂阻止新的附件引用提交，等待在途上传最终提交完成，记下数据库时间及WAL位置T，列出T之前所有已引用对象和哈希，然后恢复正常提交。后台复制这些不可变对象到B和C；确认对象全部到位、哈希正确且WAL已归档覆盖T后，才把T发布为“可恢复检查点”。复制失败时，检查点保持上一版并告警。')
p('单纯B数据库追到最新并不代表附件也追到最新。故障切换时先选择已验证的联合检查点；若B包含检查点之后缺附件的提交，必须补齐对象后再开放，或从C在B恢复到上一个联合检查点。不能把缺档伪装成正常AVAILABLE状态。B的较快RTO以该一致性验证通过为前提。')
p('夜间执行一份完整数据库引用清单校验；按备份批次记录软件版本、数据库备份ID、可连续回放WAL范围、附件清单和检查点。时间点恢复优先选择误操作之前且附件覆盖的T；选择更早T时，仓库允许存在未被该数据库版本引用的额外对象。')
h('7.5 备份频率与保留策略',2)
table(['对象或任务','计划','保留与验证'],[
('数据库物理全备','每周日01时 经pgBackRest写入C','至少保留跨越最近35天的完整备份链，并保留更早一个可用锚点；由工具管理依赖'),
('数据库差异备份','其余每天01时','保留在35天可恢复窗口内所需的差异及全备；备份与WAL不手工独立删除'),
('数据库WAL归档','持续归档 archive_timeout可从60秒起步压测','连续性校验；5分钟未归档或可恢复检查点落后超过10分钟预警 超过15分钟故障告警'),
('B数据库复制','异步流复制 持续运行','监测接收及重放位置；复制槽不能无限占磁盘；复制失败不替代C备份'),
('附件及源文件','每5分钟复制新增不可变对象至B和C并发布联合检查点','C使用restic版本化加密备份；恢复窗口内引用文件均可找回'),
('配置与程序清单','每日备份 发布及配置变更后立即加备','保留最近35天日版本及最近12个月发布和关账版本；密钥单独封存'),
('异地离线轮换','每日将已验证的联合备份复制到轮换盘 完成后断开并移至异地','记录交接时间和哈希；最近7个日版本及至少12个月度归档，容量不足时增盘'),
('月结归档','每月关账后保存业务快照及独立恢复包','至少覆盖最近12个月的可恢复版本；长期财务档案保存年限由主体适用规定和财务确定'),
('每日自动验证','读取备份清单 校验归档连续性 核对对象计数及样本哈希','保留日志；备份任务成功不等于恢复验证通过'),
('恢复演练','上线前全量演练 每月抽样恢复 每季度完整故障演练','记录实测RPO RTO和金额凭证附件结果；任何重要升级后追加演练')],[3.4,6.3,7.9])
p('每个月度归档包必须独立可恢复：包含全备或完整依赖链、使备份一致所需WAL、匹配的附件快照、版本和密钥恢复说明。不能只长期保留差异备份文件名。每个需要恢复任意时点的区间，还必须保留该区间连续WAL；月度归档包默认恢复到关账检查点，不承诺12个月内任意秒均可恢复。')
p('日常在线备份与生产数据至少分成三份：A生产、C独立备份、D异地离线；B额外提供待机复制。复制会把误删或错误更新一起传播，因此B用于缩短中断，C和D用于找回历史。RAID、同机快照或单一NAS均不能单独构成完整容灾。')
h('7.6 恢复目标与适用边界',2)
table(['故障情形','建议目标','前提及处理'],[
('A主机故障 B与C正常','整体RPO不超过15分钟；RTO不超过1小时','B数据和附件一致则手工提升；否则从C恢复至联合检查点 可能改用4小时RTO'),
('误删 错记账 逻辑损坏','技术回退目标选误操作前；隔离恢复RTO不超过4小时','优先隔离恢复并抽取受影响记录；已合法记账业务优先更正，整库回退需列出回退窗口所有后续合法交易并补录'),
('单附件误删或损坏','RPO不超过15分钟；已归档历史附件不丢失；RTO不超过2小时','按对象哈希从C找回，校验后恢复；未完成联合检查点的最新上传仍存在丢失窗口'),
('A与B故障 C完整','RPO不超过15分钟；RTO不超过4小时','有可用替代主机和离线软件包；按C最近可恢复检查点重建'),
('整处办公地点损毁','基线RPO不超过48小时；RTO不超过2个工作日','使用最近已送达异地的离线介质；需要异地机房或预留替代设备，未完成交接就不算异地副本'),
('勒索或账号大范围泄露','RPO取最近确认干净的备份；RTO由取证与重建实测','隔离生产及在线备份；优先离线副本；轮换密钥账号。潜伏时间未知，不能保证15分钟回退目标')],[3.7,6.2,7.7])
p('若要求整处办公地点损毁也只丢15分钟数据，必须增加不同地点的在线备份节点、稳定加密链路以及附件联合检查点复制；这不属于每日离线轮换可达到的能力。若要求已提交交易零丢失，还需要同步复制和配套文件一致性策略，并接受同步节点故障可能阻塞写入的取舍。')
h('7.7 主机故障切换步骤',2)
for t in [
'1 监控告警后登记事故时刻和影响范围。运维确认A失效，财务负责人暂停出款和重复录入。',
'2 将A断网或停机并封锁旧写入口，确认旧主库无法继续接受业务写入。仅切DNS不足以防止双主。',
'3 检查B数据库复制位置、连续WAL、最后联合检查点和附件哈希。选择可证明完整的恢复边界。异步复制仍可能丢失最后少量已提交交易，需在事故记录中列明。',
'4 满足一致性时提升B为主库并启动对应版本应用；不满足则先从C做受控恢复。临时不运行重复后台任务。',
'5 验证用户和权限、银行现金余额、发票待付、船东待收、凭证借贷、附件打开、单据序号和核销记录。财务签字确认后更新内部访问入口。',
'6 恢复一套正常后台任务，追查恢复窗口中可能丢失或重复的支付和导入；网银和银行对账单再次核实。记录实际RPO RTO。',
'7 A修复后从新主B重新建立从库，确认复制追平再计划回切。旧A的数据目录不得直接接回作为可写主库。'
]: p(t)
h('7.8 误操作或数据损坏恢复步骤',2)
p('先保全当前生产数据和审计证据；在隔离网络用匹配版本安装包恢复基础备份，再连续回放WAL至误操作前。核对时间线、时区和目标T，恢复对应附件对象。隔离环境停用发款、邮件和自动任务接口，避免演练或恢复副本产生外部动作。')
p('优先比较受影响记录并通过批准的更正流程修复生产，不直接把整库倒回。确需整库回退时，财务负责人批准恢复边界，列出T之后的所有合法付款、收款和新单，保留当前库副本，执行回退后按银行事实逐项重建；不能因技术回滚让真实支付消失。')
h('7.9 容量与告警如何落地',2)
p('备份容量按实测压缩和增量测算：需要容量 ≈ 保留全备数×压缩后全备大小 + 保留期内差异备份合计 + 日均WAL量×在线天数 + 附件历史版本量 + 月度归档包量，再加30%余量。示例：6份20GB全备、差异合计140GB、2GB每日WAL×35天、附件历史600GB、月包300GB，合计1230GB；加30%约1599GB。示例仅用于理解，不能直接据此采购。')
p('磁盘使用超过70%预警、85%升级；联合检查点落后10分钟预警、15分钟故障；数据库复制持续落后60秒、备份失败、WAL归档中断、校验失败、证书少于30天到期、离线介质超过36小时未完成异地交接均告警。告警发到公司现有内部通知渠道，并记录处理人；即使外网中断也应有本地声光、仪表盘或电话值守替代。')
p('每季度至少演练一次“生产主机完全不可用”，一次“误操作时间点恢复”，一次“从异地盘重建且能解密附件”。记录恢复耗时、最大数据窗口、哈希校验、金额核对、失败原因和整改闭环。未达目标时调整资源、频率或恢复流程，不把计划目标当作已实现承诺。')

h('8 分阶段实施与验收')
h('8.1 实施顺序与交付物',2)
table(['阶段','具体工作与交付物','进入下一阶段条件'],[
('0 口径及资料确认','确认主体 船名 币种 期初 会计政策及责任方；补齐银行对账单和凭证','关键身份和金额范围明确；待办有负责人与处理方式'),
('1 数据模型和样例','完成表结构 API原型 业务流程 字典和10组真实样例；建测试环境','采购 发票 付款 回款 现金和跨船示例能逐笔解释'),
('2 首期管理子账','实施M01至M09基础功能 M11期间与管理报表 M12权限与迁移','部分核销 资金分配 状态 期间 错误处理和权限验证通过'),
('3 迁移与并行核对','先暂存去重 再业务复核 导入待补字段；并行运行至少一个完整月结周期','历史兼容值一致；正式口径差异有桥接说明；无未处理关键异常进入记账'),
('4 局域网投产','配置A B C D；上线前执行完整恢复及账号切换演练；财务培训','备份可解密 数据和附件一致 权限分离 恢复目标通过实测'),
('5 扩展功能','真实盘点到位后启用M10；完善复杂估算和合同分成；满足前置条件后正式总账','规则批准 期初确认 科目映射和凭证验证完成')],[3.4,8.8,5.4])
p('具体工期和人员安排应在第0阶段确定，按交付物评估，不以固定天数代替范围确认。首期不承诺工资计算、税务申报、固定资产折旧或逐物料数量库存；这些事项可录入经确认的费用和凭证，需要完整专门业务流程时另行设计对应模块。')
h('8.2 历史迁移步骤',2)
for t in [
'1 原始文件冻结：校验原Excel哈希，保留只读副本；解析原值 原公式 缓存和单元格位置。历史快照使用2026年10月14日，方案日期与财务快照日分开。',
'2 字段归一：同名供应商 船名 日期 科目编码和汇率方向规范化；不明确币种不得按汇率1猜测USD。保留原始文本及转换规则。',
'3 重复候选复核：比对总清单和分类页订单 摘要 金额 科目及状态；支持金额分摊和多坐标映射，不能只按订单号自动去重。',
'4 业务分类：区分真实费用 转账 押金 借款 跨船代付 退款 船东直接支付；空状态与有意排除分别处理。',
'5 对账和补档：逐笔匹配银行和现金；附件文件名先建MISSING占位；原表结算统计不直接当已收回款。',
'6 分层入库：原始层完整保留；兼容层可复现旧指标；正式层只导入审核通过的事项。每次重跑导入按文件及转换版本幂等。',
'7 结账确认：导入真实期初和未结应付应收，财务审核凭证；完成余额核对和恢复演练后再启用生产写入。'
]: p(t)
h('8.3 必须通过的业务验收',2)
table(['场景','验收结果'],[
('原表兼容快照','实收371554.52；原费用398259.751018376；银行92779.85；CTM29908；保留资金差-2790.581998376'),
('分类和明细一致','五类汇总正确；485与350互换及522.51状态冲突候选逐笔确认 不靠合计抵销'),
('部分付款和并发','1000分400和600支付；两人同时核销也不能超1000；重复请求不能重复扣余额'),
('部分回款','结算1020收400再收620；待收依次620和0；同一到账不得重复核销'),
('现金划拨','银行转现金1000不增加费用；实际消费100只计一次费用；手续费独立'),
('退款与手续费','9500本金和20手续费返还，银行+9520；原付款及核销相应冲回 不新增收入'),
('计划拨款','5000×31=155000；未到账不增加资金；部分到账只更新对应实收'),
('新增坞修','本金1000加费20仅计1020；自动包含到分类核对，不重复加费'),
('截止日期和关账','截止日后收支排除；历史快照不随当天日期改变；已关账禁止覆盖金额'),
('多币种与舍入','1000 CNY×0.149=149 USD；原币保留；0.004及0.005 舍入 退款 空值与0分别验证'),
('跨船代付和船东付','出款账户与受益船不同仍可追溯；船东直接付不计公司银行支出'),
('库存和估算','100+300−80=320；盘点未知不生成耗用；除法汇率 低值0及中值缺失不混同'),
('权限和附件','越权访问被后端拒绝；不能批准自己付款；缺档可追踪；上传后哈希可校验'),
('正式记账','凭证借贷相等；费用不在应计和到票两次确认；损益及权益不以差额强行平衡'),
('恢复和容灾','断开A后从B或C恢复；误操作前PITR；异地介质可解密；金额 核销 附件 权限均验证')],[4.2,13.4])
h('8.4 建设前必须决定的口径',2)
table(['决策事项','谁确认','不确认会影响什么'],[
('公司法定名称 船名别名及账户持有主体','业务负责人 财务','主数据归属 跨船跨主体往来及凭证'),
('每种拨款是代管款 收入 借款还是报销','财务 合同负责人','收入 应收 应付及资金报表'),
('验收应计还是到票确认费用','财务负责人','费用截止期 应计转回及正式凭证'),
('税额 可抵扣税种 正式科目及期初余额','财务负责人','法定报表和应付账面值'),
('船东可报销范围 是否包含手续费','合同负责人 财务','结算金额和公司待收'),
('原表缺状态 币种不明及未解释差额','业务经办 财务','历史迁移质量及是否进入正式账'),
('日费率 天数规则 汇率源及舍入','业务负责人 财务','资金计划 日均成本 金额尾差'),
('库存金额盘点还是需要物料数量','船上负责人 财务','M10范围与后续物料子模块'),
('异地恢复目标与可投入设备','管理层 运维 财务','选择每日离线还是异地在线复制及实测RTO')],[7.5,3.8,6.3])
p('这些事项不妨碍先确认模块和界面，但在启用正式记账、结算或相关计算前必须批准。对暂时无法补齐的历史资料，应标记历史保留和质量状态，不补造审批过程、发票号或真实附件。')

h('9 完整数据表与字段字典')
h('9.1 字段类型及通用约束',2)
p('下列为完整逻辑字段设计，含通用字段和每表专用字段。物理建表时还需把本章约束落实为数据库约束、索引和服务端校验。框架自带的会话及数据库迁移版本表属于技术内部实现，不承载财务规则，按所锁定框架维护。')
table(['缩写','建议物理类型','说明'],[
('ID','UUID','全局稳定ID；PK为主键，FK为外键'),
('TXT','VARCHAR或TEXT','代码一般不超过64字符；名称200；备注和原公式使用TEXT'),
('C3','CHAR 3','币种代码，外键引用currency.code'),
('DEC','NUMERIC 20 6','业务金额；确认或记账时按对应币种小数位量化，USD正式金额2位'),
('RAW','NUMERIC 30 12','历史复算精度；更高精度原值用原始文本和原文件保存'),
('RATE','NUMERIC 20 10','汇率和比例；汇率大于0 比例通常0至1'),
('DATE','DATE','业务日期 不含时区'),
('TS','TIMESTAMPTZ','带时区时间戳，存UTC，按业务时区显示'),
('INT BIGINT BOOL','INTEGER BIGINT BOOLEAN','整数 大整数 布尔值；负数边界按业务设置'),
('ENUM','VARCHAR加CHECK约束','受控枚举；必填值不能用任意中文或英文拼写代替'),
('JSON','JSONB','结构化参数 原值和清单；需结构验证 不保存无法执行校验的任意脚本')],[2.3,5.4,9.9])
p('“必填”指正式业务可用前必须具备；草稿可以暂缺条件字段，但提交、批准或记账时按状态校验。“可空”不表示可随意省略；有合同、发票或对账依据时应补齐。“计算”表示由服务端生成或查询派生，用户不得直接更改。全局默认金额未知为NULL，确认零为0。')
h('9.2 每张表都包含的公共字段',2)
COMMON_FIELDS=[
('id','ID 必填 PK','系统生成的主键；所有表均有，不使用Excel行号或名称作主键'),
('created_at','TS 必填','创建时间 由系统写入'),
('created_by','ID 可空 FK','创建app_user；初始化及系统任务可空，并在审计中记录actor_type'),
('updated_at','TS 必填','最近变更时间；追加型表创建后保持不变'),
('updated_by','ID 可空 FK','最近变更app_user；系统任务可空'),
('row_version','INT 必填','从1递增的乐观锁版本；追加型记录固定为1，防并发覆盖')]
table(['公共字段','类型及要求','含义'],COMMON_FIELDS,[4.2,3.3,10.1])
p('所有后续数据表都实际包含以上6个字段，表内不重复印刷。source_cell_mapping、audit_log、workflow_action、已确认核销和已记账事实按追加规则写入；修正新建记录并关联原记录，不能只更新公共updated_at覆盖旧事实。')
p('退款或冲回采取追加反向事实：payment及现金使用有符号金额；收款、用途分配、贷项和船东回款分配使用effect_sign。原有效记录不能从历史累计中简单删除，反向记录从冲回生效日参与计算；展示状态REVERSED不代表删除原金额。草稿作废不计财务影响。所有余额按生效日期累计原记录与反向记录，禁止同时排除原记录又减一次反向金额。')
p('跨主体或跨币种引用默认禁止，只有明确支持的付款及结算核销可分别保存两端币种和本位币金额。多态关联如entity_type加entity_id无法直接使用普通单外键，必须通过注册类型、事务校验和定期完整性扫描保证存在性；如果落地团队需要完全数据库级约束，可改为各业务专用关联表。')
num=0
for mid,title,*_ in MODULES:
    es=[e for e in SCHEMA if e['module']==mid]
    if not es: continue
    h(mid+' '+title+'字段',2)
    for e in es:
        num+=1
        h(f'表{num:02d} {e["title"]}',3)
        p('表名：'+e['name']+'。'+e['meaning'])
        table(['字段','类型及要求','中文说明 来源与计算规则'],e['fields'],[4.7,3.0,9.9])
        if e['constraint']: p('校验与控制：'+e['constraint'])
        p('公共字段：id、created_at、created_by、updated_at、updated_by、row_version，详见9.2。')

h('9.3 必须建立的唯一约束和索引',2)
table(['对象','约束或索引'],[
('业务编号','organization_id加单据编号联合唯一；编号不能用最大行号加1无锁生成'),
('科目与币种','account_code的namespace加code唯一；gl_account主体加code唯一；currency.code唯一'),
('行号和关系','event_id加line_no invoice_id加line_no entry_id加line_no分别唯一；角色权限对唯一'),
('库存月结','主体 船 月份 库存类型 币种联合唯一'),
('导入和幂等','source_file主体加sha256防重复批次；journal_entry.generation_key全局唯一；审批request_id与动作幂等'),
('流水防重','银行真实external_reference在账户内按银行规范唯一；无可靠编号使用fingerprint候选复核'),
('报表快照','组织 船 报表码 期间 截止日 规则及snapshot_version唯一；空船维度也必须落实唯一性'),
('常用查询','按主体 船 posting_date及状态建组合索引；核销按发票行 付款及生效日；银行按账户 日期 排序键'),
('外键和期间','外键查询列建索引；已用主数据禁止删除；同主体会计期间及生效参数区间不得重叠'),
('核销事务','锁定付款及发票可用余额；保存分配与审计同一事务完成；数据库CHECK验证正负 0及枚举边界')],[4.4,13.2])

h('10 原工作表到新模块的迁移对照')
table(['原工作表','新模块和去向','迁移注意事项'],[
('SOFP','M11 report_snapshot report_metric','保留旧口径；股本及收益不能直接用倒算值初始化正式总账'),
('Cost Checking','M11管理费用核对','按统一业务行分组；加入坞修，不复制旧固定范围'),
('Bank Verification','M06账户 流水及bank_match','22条流水分离本金和手续费，补外部参考号及对账单'),
('Shipmanagement RCVD','M05实际收款分配 M09资金申请','9条有日期已分配资金与3条结算差额及1条计划分开核实'),
('Expenses Lists','M02 M03 M04唯一业务链','98条摘要行不是98条已确认费用；含现金补给 退款和空金额'),
('Shipmanagement Cost','M02同一台账 船管中心','58条摘要行与总表合并，识别跨船 借款和资产调整'),
('Owner Cost','M02船东中心 M08结算','11条摘要行；付款状态和船东结算状态各自保留'),
('Take-Over Cost','M02接船中心 M08结算','16条摘要行；GRN与发票取值不同及AE15错误单独审查'),
('BMA Flag Change','M02换旗中心 M08结算','14条摘要行；船东直接付保留，附件U V W含义按本页识别'),
('DD Cost','M02坞修中心','当前无业务；需测试非零本金及手续费，避免重复加费'),
('CTM Expenses List','M07现金账户 流水和划拨','22条流水分成2条入金20条支出；关联费用避免重复'),
('分析','M09估算及参数版本','63条分类模板；除法汇率，低值例外，中值未知；自由辅助区归档'),
('Account Code-2版','M01 account_code','160船管及32船员细项；各自保留命名空间'),
('伙食海关库房管理','M10金额库存月结 M07现金','7月8月零金额模板；修正跨月引用前用真实非零样例确认')],[4.5,6.3,6.8])
h('10.1 原费用表中容易遗漏的字段如何保存',2)
table(['原字段','规划去向和处理'],[
('入账日期 类型 RQA Supplier PO Nos','business_event.posting_date和description；party；purchase_order_line.requisition_no和po_no'),
('PO Current Rating PO USD','订单原币 汇率和base_amount；原表无币种时产生异常，原值保留'),
('Account Code GRN Date GRN Amount','account_code及goods_receipt.receipt_date和base_amount'),
('Invoice Amount Invoice Date Due Date','invoice及invoice_line；缺实际发票号不能用附件文件名推断'),
('C/N Accrual C/N Profit Profit','所有原值进入source_cell_mapping；确认是贷项 应计或利润后分别转业务或凭证，不先定义通用利润公式'),
('Payment Status 六个状态金额列','legacy_status和兼容F21；新核销和审批采用独立事实'),
('Bank Charge Bank Slip','实际费用与bank_transaction匹配；回单转attachment及attachment_link'),
('Pay Approval 各Attachement列','按审批 发票 贷项 补充材料建多附件关系；缺文件为MISSING'),
('船东已收 待收 船东付等','保留原文本；M08真实待收由结算与回款核销计算'),
('CTM In Out Others Food Store','转有符号现金流水；分类由业务行及库存类型确认，原值完整保留')],[5.4,12.2])

h('附录A 原表45项规则与迁移处置')
p('本附录把现有分析中R01至R45逐项转为可读说明。证据为原工作表及单元格坐标，可在只读源文件中核查。原表错误和待确认口径不直接成为新系统的自动规则。')
spec=json.loads((ROOT/'analysis/system_specification.json').read_text())
for code,title,evidence,meaning,action in spec['rules']:
    h(code+' '+title,3)
    p('原表依据：'+evidence)
    p('含义或现状：'+meaning)
    p('系统处理：'+action)

h('附录B 资料依据和技术参考')
p('业务资料以项目内source目录的原工作簿及data和analysis目录的提取分析为依据。原文件和现有数据未作改写；本方案是建设规划，实施、部署和财务政策审批分别执行。')
table(['项目资料','用途'],[
('source/MV PIONEER LAY 费用统计SOFP--20261014(1).xlsx','原始14个工作表及保存快照；SHA256见下文'),
('analysis/system_specification.json','14页角色 45条业务规则 建议实体 验收和待确认事项'),
('analysis/data_verification.json','源哈希 提取完整性 公式复算和关键汇总核对结果'),
('data/field_dictionary.csv','原工作表267条字段说明；列位与语义归一依据'),
('data/findings.csv','25项已出现 潜在 缺证及待确认发现'),
('data/business_records.csv 和 bank_transactions.csv','业务摘要 付款状态及22条银行原始流水'),
('data/all_formulas.csv 和 report_output_mapping.csv','原公式精确文本及报表坐标到指标映射'),
('data/account_codes.csv','两套192条科目编码'),
('data/duplicate_view_candidates.csv 和 amount_differences.csv','重复候选及阶段金额差异，不等于已确认重复或错误')],[8.4,9.2])
p('原工作簿SHA256：4b77e28208fa28fad5752fd532d421fe5151c992f7ee60c3c02d08d847facfdc')
p('以下为官方技术文档，核验日期2026年10月15日。备份时点恢复、异步复制和工具能力依据这些资料；频率、保留天数、资源及RPO RTO为本方案建议，须经公司确认并演练。')
SOURCES=[
('PostgreSQL 时间点恢复与连续归档','https://www.postgresql.org/docs/18/continuous-archiving.html'),
('PostgreSQL 流复制与同步复制','https://www.postgresql.org/docs/18/warm-standby.html'),
('PostgreSQL 版本支持政策','https://www.postgresql.org/support/versioning/'),
('pgBackRest 备份与恢复操作指南','https://pgbackrest.org/user-guide.html'),
('pgBackRest 命令及保留策略','https://pgbackrest.org/command.html'),
('restic 追加写入仓库及清理权限说明','https://restic.readthedocs.io/en/v0.17.0/060_forget.html'),
('Django 支持版本与维护期限','https://www.djangoproject.com/download/'),
('Django Python版本兼容说明','https://docs.djangoproject.com/en/dev/faq/install/'),
('飞书 上传本地文件及导入为云文档','https://www.feishu.cn/hc/en-US/articles/360033241654-upload-or-import-local-files-and-folders')]
for title,url in SOURCES: link(title,url)
h('附录C 导入飞书文档的方法')
p('在飞书云文档中选择上传或导入，选择导入为新文档并选取本Word文件；也可以先上传docx，再通过菜单转换为飞书文档。飞书官方支持docx导入。导入后检查标题层级、表格列宽和分页附近的内容；Word页码不作为飞书文档中的章节定位依据。')
p('本文使用原生Word标题、段落和表格，公式为可编辑文本，没有HTML页面或需联网运行的嵌入脚本。大部分阅读先看第1至3章；开发评审再查第4至9章和原表规则附录。')

target=OUT/'船舶管理财务系统建设方案与数据字典.docx'
# 移除基础Word模板中可能继承的段落装饰线及字体主题覆盖。
for element in [doc.styles.element,doc.element]:
    for b in list(element.iter(qn('w:pBdr'))): b.getparent().remove(b)
for st in doc.styles:
    if st.type==1 and st.element.rPr is not None and st.element.rPr.rFonts is not None:
        fonts=st.element.rPr.rFonts
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+key),None)
doc.save(target)
# 仅供内部完整性检查与以后修订的结构清单，不作为额外交付件。
(ROOT/'analysis/plan_schema.json').write_text(json.dumps(SCHEMA,ensure_ascii=False,indent=2))
print(json.dumps({'path':str(target),'tables':len(SCHEMA),'specialized_fields':sum(len(e['fields']) for e in SCHEMA),'common_fields_per_table':6,'word_tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))




