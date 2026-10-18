from pathlib import Path
import openpyxl, json, collections, datetime, shutil, hashlib
from openpyxl.formula.tokenizer import Tokenizer

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('/Users/stone/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/s269119257_97a6/msg/file/2026-10/MV PIONEER LAY 费用统计SOFP--20261014(1).xlsx')
w=openpyxl.load_workbook(SOURCE, data_only=False)
v=openpyxl.load_workbook(SOURCE, data_only=True)
if not (ROOT/'source'/SOURCE.name).exists():
    shutil.copy2(SOURCE, ROOT/'source'/SOURCE.name)
def ser(x):
    return x.isoformat() if isinstance(x,(datetime.datetime,datetime.date,datetime.time)) else x
def raw(c):
    return c.value.text if hasattr(c.value,'text') else c.value
cells=[]; sheets=[]; rows=[]; functions=collections.Counter()
for s in w:
    nonempty=[c for c in s._cells.values() if c.value is not None]
    nonempty.sort(key=lambda c:(c.row,c.column))
    sheets.append(dict(name=s.title,state=s.sheet_state,dimension=s.calculate_dimension(),nonempty=len(nonempty),formulas=sum(c.data_type=='f' for c in nonempty),merges=[str(m) for m in s.merged_cells.ranges],hidden_rows=[i for i,d in s.row_dimensions.items() if d.hidden],hidden_columns=[i for i,d in s.column_dimensions.items() if d.hidden],tables=list(s.tables),validations=[dict(range=str(d.sqref),type=d.type,formula1=d.formula1,formula2=d.formula2,allow_blank=d.allowBlank,operator=d.operator,error=d.error) for d in s.data_validations.dataValidation],freeze=str(s.freeze_panes),auto_filter=str(s.auto_filter.ref),print_area=str(s.print_area)))
    rows.append('\n### '+s.title)
    for r, cs in __import__('itertools').groupby(nonempty,key=lambda c:c.row):
        cs=list(cs)
        rows.append(' | '.join(f'{c.coordinate}: {str(raw(c)).replace(chr(10)," / ")}'+(f' [cache={v[s.title][c.coordinate].value}]' if c.data_type=='f' else '') for c in cs))
    for c in nonempty:
        val=v[s.title][c.coordinate].value
        isf=c.data_type=='f'
        if isf:
            functions.update(t.value[:-1].upper() for t in Tokenizer(raw(c)).items if t.type=='FUNC' and t.subtype=='OPEN')
        cells.append(dict(sheet=s.title,cell=c.coordinate,row=c.row,column=c.column,kind='formula' if isf else 'literal',data_type=c.data_type,value=ser(val),raw_value=ser(raw(c)),formula=raw(c) if isf else '',array_range=getattr(c.value,'ref',''),number_format=c.number_format,is_date=c.is_date,comment=c.comment.text if c.comment else '',hyperlink=c.hyperlink.target if c.hyperlink else '',hidden_row=s.row_dimensions[c.row].hidden,hidden_column=s.column_dimensions[c.column_letter].hidden))
result=dict(source_name=SOURCE.name,source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),sheets=sheets,cells=cells,functions=functions,calculation=str(w.calculation),external_links=len(w._external_links),defined_names=[str(n) for n in w.defined_names.values()])
(ROOT/'data/workbook_raw.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'analysis/rows.txt').write_text('\n'.join(rows),encoding='utf-8')
print(json.dumps(dict(sheets=len(sheets),cells=len(cells),formulas=sum(s['formulas'] for s in sheets),functions=functions,calculation=str(w.calculation)),ensure_ascii=False,indent=2))
